import gradio as gr
import torch
import os
import asyncio
import numpy as np
import math
import cv2
import socket
import json
import tempfile
import base64
import zipfile
import re
import html
import pickle
from datetime import datetime
from io import BytesIO
from PIL import Image

from glls import paths as glls_paths
from glls.gradio_compat import patch_gradio_compat
from glls.weld.frontend import (
    WELD_CATEGORY,
    WELD_DATASET,
    WELD_TASK_PREFIX,
    WeldFrontendRuntime,
)
from glls.runtime_config import (
    abound_dataset_weight_paths,
    default_adaptclip_checkpoint,
    normalize_dataset_key,
    parse_k_shot,
    published_localizer_name,
    resolve_adaptclip_checkpoint,
    resolve_localizer_name,
    runtime_weight_config,
    runtime_weight_summary,
)

from transformers import (
    AutoProcessor,
    Qwen2_5_VLForConditionalGeneration,
    LlavaOnevisionForConditionalGeneration
)

# ==========================================
# 0. Environment & Setup
# ==========================================
os.environ["VLLM_WORKER_MULTIPROC_METHOD"] = "spawn"
GLLS_PROJECT_ROOT = glls_paths.project_root()
GLLS_DATA_ROOT = glls_paths.data_root()

def _default_existing_path(configured_path, project_relative_path):
    project_path = os.path.join(GLLS_PROJECT_ROOT, project_relative_path)
    return project_path if os.path.exists(project_path) else configured_path

DEFAULT_MMAD_ROOT = _default_existing_path(glls_paths.dataset_root(), "datasets/MMAD")
DEFAULT_QA_ROOT = _default_existing_path(glls_paths.qa_root(), "qa_collection")
DEFAULT_GRAPH_ROOT = glls_paths.graph_cache_root()
DEFAULT_SAM3_PATH = glls_paths.sam3_path()
DEFAULT_VLM_MODEL_PATH = glls_paths.vlm_model_path()


def _default_adaptclip_checkpoint(dataset_name):
    return default_adaptclip_checkpoint(dataset_name)


def _resolve_adaptclip_checkpoint(path_value, dataset_name):
    return resolve_adaptclip_checkpoint(path_value, dataset_name)


DEFAULT_ADAPTCLIP_CHECKPOINT_PATH = _default_adaptclip_checkpoint("mvtec")


def _default_gpu_id():
    configured = os.environ.get("GLLS_DEFAULT_GPU") or os.environ.get("GLLS_GPU_ID")
    if configured:
        return configured.split(",")[0].strip()
    return "0"


def _gpu_choices():
    try:
        if torch.cuda.is_available():
            count = torch.cuda.device_count()
            if count > 0:
                return [str(i) for i in range(count)]
    except Exception:
        pass
    return [str(i) for i in range(8)]


def _normalize_gpu_id(gpu_id):
    gpu_text = str(gpu_id).strip()
    if not gpu_text:
        return 0
    return int(gpu_text)


def _normalize_dataset_name(dataset_name):
    dataset_key = normalize_dataset_key(dataset_name)
    return "visa" if dataset_key == "visa" else "mvtec"


def _parse_k_shot(k_shot):
    return parse_k_shot(k_shot)


def _frontend_localizer_name(dataset_name, k_shot):
    return published_localizer_name(dataset_name, k_shot)


def _normalize_localizer_choice(localizer_choice, dataset_name, k_shot):
    return resolve_localizer_name(localizer_choice, dataset_name, k_shot)


def _abound_dataset_weight_paths(save_path, dataset_name):
    return abound_dataset_weight_paths(save_path, dataset_name)


def _runtime_weight_config(dataset_name, localizer_choice, k_shot, adaptclip_ckpt_path):
    return runtime_weight_config(dataset_name, localizer_choice, k_shot, adaptclip_ckpt_path)


def _runtime_weight_summary(config):
    return runtime_weight_summary(config)


def _runtime_config_error_html(message):
    return (
        "<div class='method-process-card'>"
        "<h3>Runtime configuration required</h3>"
        f"<p>{_html(message)}</p>"
        "</div>"
    )

# Try importing vLLM
try:
    from vllm import LLM, SamplingParams
    HAS_VLLM = True
    print("vLLM library detected.")
except ImportError:
    HAS_VLLM = False
    print("vLLM not found. Acceleration disabled.")

# Import Project Modules
try:
    from glls.models.mcts_sam import MCTSQuestionSample
    from glls.models.localizer import ABounD_Localizer, AdaptCLIP_Localizer
    from glls.qa_trace import binary_anomaly_decision_override, summarize_method_participation
    from glls.rag.agent import SimInspecAgent
    from glls.seg.sam3_engine import Sam3Engine
    # Import Loaders
    from glls.data.dataset_loader import DSMVTecDatasetLoader, VisaDatasetLoader 
except ImportError as e:
    print(f"Project module import failed: {e}")
    print("Ensure you are running this from the project root.")

# Qwen3 Support
try:
    from transformers import Qwen3VLForConditionalGeneration
    HAS_QWEN3 = True
except ImportError:
    HAS_QWEN3 = False

# ==========================================
# 1. Utility Functions
# ==========================================
def find_free_port(start_port):
    port = start_port
    while True:
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            if s.connect_ex(('localhost', port)) != 0:
                return port
            port += 1

def image_to_base64(image):
    buffered = BytesIO()
    image.save(buffered, format="PNG")
    return base64.b64encode(buffered.getvalue()).decode('utf-8')

def _as_text(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    try:
        return str(value)
    except Exception:
        return ""

def _coerce_options(value):
    if isinstance(value, dict):
        return value
    if not value:
        return {}
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
            return parsed if isinstance(parsed, dict) else {}
        except json.JSONDecodeError:
            return {}
    return {}

def _extract_path(item):
    if item is None:
        return None
    if isinstance(item, str):
        return item
    if isinstance(item, dict):
        return item.get("path") or item.get("name")
    if isinstance(item, (list, tuple)) and len(item) > 0:
        return _extract_path(item[0])
    if hasattr(item, "name"):
        return getattr(item, "name")
    return None

def _pil_from_any(image_like):
    if image_like is None:
        return None
    if isinstance(image_like, Image.Image):
        return image_like
    if isinstance(image_like, np.ndarray):
        try:
            return Image.fromarray(image_like)
        except Exception:
            return None
    path = _extract_path(image_like)
    if path and os.path.exists(path):
        try:
            return Image.open(path).convert("RGB")
        except Exception:
            return None
    return None


def _image_data_uri(image_like, max_px=900):
    image = _pil_from_any(image_like)
    if image is None:
        return ""
    image = image.convert("RGB").copy()
    image.thumbnail((int(max_px), int(max_px)))
    buffer = BytesIO()
    image.save(buffer, format="JPEG", quality=86)
    return "data:image/jpeg;base64," + base64.b64encode(buffer.getvalue()).decode("utf-8")


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}


def find_normal_support_images(dataset_root, dataset_name, category, limit):
    if int(limit or 0) <= 0:
        return []

    dataset_key = str(dataset_name or "").strip().lower()
    category = str(category)
    if dataset_key == "mvtec":
        candidate_dirs = [
            os.path.join(dataset_root, category, "image", "good"),
            os.path.join(dataset_root, category, "train", "good"),
            os.path.join(dataset_root, category, "test", "good"),
        ]
    else:
        candidate_dirs = [
            os.path.join(dataset_root, category, "train", "good"),
            os.path.join(dataset_root, category, "test", "good"),
            os.path.join(dataset_root, category, "image", "good"),
        ]

    images = []
    for directory in candidate_dirs:
        if not os.path.isdir(directory):
            continue
        for name in sorted(os.listdir(directory)):
            path = os.path.join(directory, name)
            if os.path.isfile(path) and os.path.splitext(name.lower())[1] in IMAGE_EXTENSIONS:
                images.append(path)
        if images:
            break
    return images[: int(limit)]


def _zip_write_text(zf, arcname, text):
    data = _as_text(text).encode("utf-8", errors="replace")
    zf.writestr(arcname, data)

def _zip_write_pil_png(zf, arcname, pil_img):
    pil_img = _pil_from_any(pil_img)
    if pil_img is None:
        return False
    bio = BytesIO()
    pil_img.save(bio, format="PNG")
    zf.writestr(arcname, bio.getvalue())
    return True

def _zip_write_path(zf, src_path, arcname):
    if not src_path or not os.path.exists(src_path):
        return False
    try:
        zf.write(src_path, arcname=arcname)
        return True
    except Exception:
        return False

def _normalize_gallery_items(gallery_value):
    if not gallery_value:
        return []
    if not isinstance(gallery_value, list):
        return []
    normalized = []
    for item in gallery_value:
        if isinstance(item, (list, tuple)) and len(item) >= 2:
            normalized.append({"image": item[0], "label": _as_text(item[1])})
        else:
            normalized.append({"image": item, "label": ""})
    return normalized

def _maybe_generate_docx_report(
    *,
    out_docx_path,
    timestamp_iso,
    category,
    task_type,
    question,
    ground_truth,
    options,
    rag_text,
    final_diagnosis_md,
    phase2_prompt,
    phase1_prompt_md,
    logs,
    input_image_pil,
    heatmap_pil,
    logic_view_pil,
    local_refinement_paths,
    rag_image_paths,
):
    try:
        from docx import Document  # type: ignore
        from docx.shared import Inches  # type: ignore
    except Exception as e:
        return False, f"python-docx not available: {e}"

    def _add_picture(doc, pil_img=None, path=None, caption=None, width_inches=6.0):
        if caption:
            doc.add_paragraph(_as_text(caption))
        if path and os.path.exists(path):
            doc.add_picture(path, width=Inches(width_inches))
            return True
        pil_img = _pil_from_any(pil_img)
        if pil_img is None:
            return False
        tmp_png = os.path.join(os.path.dirname(out_docx_path), f"_tmp_{os.urandom(6).hex()}.png")
        pil_img.save(tmp_png)
        doc.add_picture(tmp_png, width=Inches(width_inches))
        return True

    try:
        doc = Document()
        doc.add_heading("GLLS QA Method Run Report", level=0)
        doc.add_paragraph(f"Timestamp: {timestamp_iso}")
        if category:
            doc.add_paragraph(f"Category: {category}")
        if task_type:
            doc.add_paragraph(f"Task Type: {task_type}")

        doc.add_heading("Question", level=1)
        doc.add_paragraph(_as_text(question))
        if ground_truth:
            doc.add_paragraph(f"Ground Truth: {ground_truth}")
        if options:
            doc.add_paragraph("Options (JSON):")
            doc.add_paragraph(json.dumps(options, ensure_ascii=False, indent=2))

        doc.add_heading("Input Image", level=1)
        _add_picture(doc, pil_img=input_image_pil, caption="Input View")

        doc.add_heading("RAG Knowledge", level=1)
        doc.add_heading("Text Rules", level=2)
        doc.add_paragraph(_as_text(rag_text))
        if rag_image_paths:
            doc.add_heading("Visual Standards", level=2)
            for i, p in enumerate(rag_image_paths, start=1):
                _add_picture(doc, path=p, caption=f"RAG Image {i}: {os.path.basename(p)}", width_inches=5.5)

        doc.add_heading("Global Anomaly Heatmap", level=1)
        _add_picture(doc, pil_img=heatmap_pil, caption="Global Anomaly Heatmap (Overlay)")

        doc.add_heading("Local Refinement", level=1)
        for i, p in enumerate(local_refinement_paths or [], start=1):
            _add_picture(doc, path=p, caption=f"Refinement View {i}: {os.path.basename(p)}", width_inches=5.5)

        doc.add_heading("Logic Debug", level=1)
        _add_picture(doc, pil_img=logic_view_pil, caption="Logic Debug View")
        doc.add_heading("Phase 1 Prompt", level=2)
        doc.add_paragraph(_as_text(phase1_prompt_md))

        doc.add_heading("Final", level=1)
        doc.add_heading("Final Prompt (Phase 2)", level=2)
        doc.add_paragraph(_as_text(phase2_prompt))
        doc.add_heading("Final Diagnosis", level=2)
        doc.add_paragraph(_as_text(final_diagnosis_md))

        doc.add_heading("Logs", level=1)
        doc.add_paragraph(_as_text(logs))

        doc.save(out_docx_path)
        return True, "ok"
    except Exception as e:
        return False, f"docx generation failed: {e}"

def export_current_run_bundle(
    input_image,
    question,
    ground_truth,
    category,
    task_type,
    options,
    rag_gallery,
    rag_text,
    heatmap_image,
    sam3_gallery,
    focus_files,
    final_diagnosis_md,
    phase2_prompt,
    logic_view_image,
    phase1_prompt_md,
    logs_text,
    include_docx,
):
    final_diagnosis_md = _as_text(final_diagnosis_md).strip()
    if not final_diagnosis_md or "Waiting for results" in final_diagnosis_md:
        return None, "No completed run to export yet."
    options = _coerce_options(options)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    timestamp_iso = datetime.now().isoformat(timespec="seconds")
    export_root = tempfile.mkdtemp(prefix="glls_qa_export_")
    zip_path = os.path.join(export_root, f"glls_qa_run_{timestamp}.zip")
    top_dir = f"glls_qa_run_{timestamp}"

    rag_image_paths = []
    if rag_gallery:
        for item in rag_gallery:
            p = _extract_path(item)
            if p and os.path.exists(p):
                rag_image_paths.append(p)

    local_refinement_paths = []
    if focus_files:
        for item in focus_files:
            p = _extract_path(item)
            if p and os.path.exists(p):
                local_refinement_paths.append(p)

    sam3_items = _normalize_gallery_items(sam3_gallery)

    manifest = {
        "export_version": 1,
        "timestamp": timestamp_iso,
        "category": _as_text(category),
        "task_type": _as_text(task_type),
        "question": _as_text(question),
        "ground_truth": _as_text(ground_truth),
        "options": options,
        "has_docx": False,
    }

    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        _zip_write_text(
            zf,
            f"{top_dir}/README.txt",
            "This ZIP contains one GLLS QA method run export.\n"
            "Folders:\n"
            "  question/  rag/  global/  local_refinement/  logic_debug/  final/  logs/\n",
        )

        # Question / metadata
        _zip_write_text(zf, f"{top_dir}/question/question.txt", question)
        _zip_write_text(zf, f"{top_dir}/question/ground_truth.txt", ground_truth)
        if isinstance(options, dict) and options:
            _zip_write_text(zf, f"{top_dir}/question/options.json", json.dumps(options, ensure_ascii=False, indent=2))
        if _pil_from_any(input_image) is not None:
            _zip_write_pil_png(zf, f"{top_dir}/question/input_view.png", input_image)

        # RAG
        _zip_write_text(zf, f"{top_dir}/rag/text_rules.txt", rag_text)
        for i, p in enumerate(rag_image_paths, start=1):
            base = os.path.basename(p) or f"rag_{i}.png"
            _zip_write_path(zf, p, f"{top_dir}/rag/visual_standards/{i:02d}_{base}")

        # Global
        _zip_write_pil_png(zf, f"{top_dir}/global/global_anomaly_heatmap.png", heatmap_image)

        # Local refinement (files)
        for i, p in enumerate(local_refinement_paths, start=1):
            base = os.path.basename(p) or f"refinement_{i}.png"
            _zip_write_path(zf, p, f"{top_dir}/local_refinement/{i:02d}_{base}")

        # Local refinement (gallery snapshots)
        if sam3_items:
            sam3_index = []
            for i, it in enumerate(sam3_items, start=1):
                label = _as_text(it.get("label", ""))
                arc = f"{top_dir}/local_refinement/gallery_{i:02d}.png"
                ok = _zip_write_pil_png(zf, arc, it.get("image"))
                if ok:
                    sam3_index.append({"file": f"local_refinement/gallery_{i:02d}.png", "label": label})
            if sam3_index:
                _zip_write_text(zf, f"{top_dir}/local_refinement/gallery_index.json", json.dumps(sam3_index, ensure_ascii=False, indent=2))

        # Logic debug
        _zip_write_pil_png(zf, f"{top_dir}/logic_debug/view_image.png", logic_view_image)
        _zip_write_text(zf, f"{top_dir}/logic_debug/phase1_prompt.md", phase1_prompt_md)

        # Final
        _zip_write_text(zf, f"{top_dir}/final/final_prompt.txt", phase2_prompt)
        _zip_write_text(zf, f"{top_dir}/final/final_diagnosis.md", final_diagnosis_md)

        # Logs
        _zip_write_text(zf, f"{top_dir}/logs/live_logs.txt", logs_text)

        # Optional DOCX
        include_docx = bool(include_docx)
        if include_docx:
            docx_path = os.path.join(export_root, f"siminspec_run_{timestamp}.docx")
            ok, msg = _maybe_generate_docx_report(
                out_docx_path=docx_path,
                timestamp_iso=timestamp_iso,
                category=category,
                task_type=task_type,
                question=question,
                ground_truth=ground_truth,
                options=options if isinstance(options, dict) else {},
                rag_text=rag_text,
                final_diagnosis_md=final_diagnosis_md,
                phase2_prompt=phase2_prompt,
                phase1_prompt_md=phase1_prompt_md,
                logs=logs_text,
                input_image_pil=input_image,
                heatmap_pil=heatmap_image,
                logic_view_pil=logic_view_image,
                local_refinement_paths=local_refinement_paths,
                rag_image_paths=rag_image_paths,
            )
            if ok and os.path.exists(docx_path):
                manifest["has_docx"] = True
                _zip_write_path(zf, docx_path, f"{top_dir}/report.docx")
            else:
                _zip_write_text(zf, f"{top_dir}/report_docx_error.txt", msg)

        # Manifest (write once, after optional DOCX)
        _zip_write_text(zf, f"{top_dir}/manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))

    docx_note = " (with report.docx)" if manifest.get("has_docx") else ""
    return zip_path, f"Export ready: {os.path.basename(zip_path)}{docx_note}"

# ==========================================
# 2. Dataset Manager (INTELLIGENT PATH FIX)
# ==========================================
class DatasetManager:
    def __init__(self, root_path, qa_root_path=None, dataset_name="mvtec"):
        self.root_path = root_path
        self.qa_root_path = qa_root_path
        self.dataset_name = dataset_name
        self.loader = None
        self.loader_data_root = None
        self.loader_qa_root = None
        self.all_samples = []          
        self.subfolder_choices = []    
        self.task_type_choices = []    

    def _target_folder(self):
        dataset_folder_map = {
            "visa": "VisA",
            "mvtec": "DS-MVTec",
        }
        return dataset_folder_map.get(str(self.dataset_name).lower())

    def _resolve_dataset_root(self):
        raw_root = os.path.normpath(self.root_path)
        target_folder = self._target_folder()
        if os.path.basename(raw_root) == target_folder:
            return raw_root
        return os.path.join(raw_root, target_folder)

    def _resolve_qa_root(self, final_data_root):
        target_folder = self._target_folder()
        raw_qa_root = os.path.normpath(self.qa_root_path or "")
        if not raw_qa_root:
            return final_data_root

        if os.path.basename(raw_qa_root) == target_folder:
            return raw_qa_root

        dataset_specific_qa = os.path.join(raw_qa_root, target_folder)
        if os.path.exists(dataset_specific_qa):
            return dataset_specific_qa

        return raw_qa_root
        
    def load_subclass(self, category):
        """Loads category data with auto-path correction."""
        
        # --- 1. Intelligent Path Resolution ---
        target_folder = self._target_folder()

        if not target_folder:
            return f"Unknown dataset type: {self.dataset_name}"

        final_root = self._resolve_dataset_root()
        final_qa_root = self._resolve_qa_root(final_root)
        if not os.path.exists(final_root):
            return f"Dataset path not found: {final_root}. Please check Dataset Root."
        if not os.path.exists(final_qa_root):
            return f"QA path not found: {final_qa_root}. Please check QA Collection Root."

        # --- 2. Loader Initialization ---
        loader_needs_init = (
            self.loader is None or
            self.loader_data_root != final_root or
            self.loader_qa_root != final_qa_root
        )
        if loader_needs_init:
            if self.dataset_name.lower() == "visa":
                print(f"Initializing VisaDatasetLoader | data={final_root} | qa={final_qa_root}")
                self.loader = VisaDatasetLoader(root_path=final_root, qa_root_path=final_qa_root)
            else:
                print(f"Initializing DSMVTecDatasetLoader | data={final_root} | qa={final_qa_root}")
                self.loader = DSMVTecDatasetLoader(root_path=final_root, qa_root_path=final_qa_root)
            self.loader_data_root = final_root
            self.loader_qa_root = final_qa_root
            
        try:
            # The loader now looks inside ".../VisA/candle" correctly
            self.all_samples = list(self.loader.parse_samples(category))
            
            subfolders = set()
            task_types = set()
            for s in self.all_samples:
                path_parts = s.get('image_path', '').replace('\\', '/').split('/')
                # Simple subfolder parsing (usually 2nd to last)
                if len(path_parts) >= 2:
                    subfolders.add(path_parts[-2])
                task_types.add(s.get('task_type', 'Unknown'))
            
            self.subfolder_choices = sorted(list(subfolders))
            self.task_type_choices = sorted(list(task_types))
            
            return (
                f"[{self.dataset_name.upper()}] {category} loaded: {len(self.all_samples)} QA rows"
            )
        except Exception as e:
            import traceback
            traceback.print_exc()
            return f"Load failed: {str(e)}"

    def filter_samples(self, subfolder=None, task_type=None, search_query=""):
        choices = []
        for idx, s in enumerate(self.all_samples):
            path_match = True
            if subfolder and subfolder != "All":
                path_match = subfolder in s.get('image_path', '')
            
            type_match = True
            if task_type and task_type != "All":
                type_match = (s.get('task_type') == task_type)
            
            img_path = s.get('image_path', '')
            file_name = os.path.basename(img_path) if img_path else "NoImg"

            search_match = True
            if search_query:
                q_text = s.get('question', '').lower()
                search_match = (search_query.lower() in q_text) or (search_query.lower() in file_name.lower())

            if path_match and type_match and search_match:
                short_q = s['question'][:40].replace("\n", " ") 
                gt = s.get('gt_answer', '?')
                choices.append(f"{idx} | [{file_name}] {short_q}... (GT: {gt})")
        
        return choices

    def get_sample_by_idx(self, global_idx):
        if 0 <= global_idx < len(self.all_samples):
            s = self.all_samples[global_idx]
            img_path = s['image_path']
            if os.path.exists(img_path):
                img = Image.open(img_path).convert("RGB")
                # Ensure options is a dict
                opts = s.get('options', {})
                if isinstance(opts, str):
                    try:
                        opts = json.loads(opts)
                    except (TypeError, json.JSONDecodeError):
                        opts = {}
                return img, s['question'], s['task_type'], s.get('gt_answer', 'Unknown'), opts
        return None, "", "", "", {}

# ==========================================
# 3. Unified VLM Inference
# ==========================================
class UnifiedVLMInference:
    def __init__(self, model_path, device, model_type="qwen2.5-vl", use_vllm=False, gpu_util=0.85):
        self.device = device 
        self.model_type = model_type.lower()
        self.model_path = model_path
        self.use_vllm = use_vllm and HAS_VLLM
        
        print(f"Loading VLM: {self.model_type} | vLLM enabled: {self.use_vllm}")

        # --- vLLM Backend ---
        if self.use_vllm:
            print(f"⚡ Initializing vLLM Engine (GPU Util: {gpu_util})...")
            try:
                self.model = LLM(
                    model=model_path,
                    trust_remote_code=True,
                    gpu_memory_utilization=gpu_util, 
                    tensor_parallel_size=1,
                    dtype="bfloat16",       
                    max_model_len=8192,     
                    enforce_eager=False,
                    disable_custom_all_reduce=True
                )
                self.sampling_params = SamplingParams(temperature=0.0, max_tokens=4096)
                print("vLLM engine ready.")
                return
            except Exception as e:
                print(f"vLLM init failed: {e}")
                print("Fallback to Transformer mode...")
                self.use_vllm = False

        # --- Transformers Backend Fallback ---
        if "qwen3" in self.model_type:
            if not HAS_QWEN3:
                raise ImportError("Qwen3 class not available.")
            self.processor = AutoProcessor.from_pretrained(model_path, trust_remote_code=True)
            self.model = Qwen3VLForConditionalGeneration.from_pretrained(
                model_path, torch_dtype="auto", device_map=device, trust_remote_code=True
            )
        elif "qwen" in self.model_type: # qwen2.5-vl
            self.processor = AutoProcessor.from_pretrained(model_path, min_pixels=256*28*28, max_pixels=1280*28*28)
            self.model = Qwen2_5_VLForConditionalGeneration.from_pretrained(
                model_path, torch_dtype=torch.bfloat16, device_map=device, attn_implementation="sdpa"
            )
        elif "llava" in self.model_type:
            self.processor = AutoProcessor.from_pretrained(model_path)
            self.model = LlavaOnevisionForConditionalGeneration.from_pretrained(
                model_path, torch_dtype=torch.float16, device_map=device
            )
        else:
            raise ValueError(f"Unsupported model type: {model_type}")
        
        if hasattr(self, 'model'):
            self.model.eval()

    def generate(self, content_list, max_tokens=4096):
        clean_content = []
        pil_images = []
        
        for item in content_list:
            if item['type'] == 'text':
                clean_content.append(item)
            elif item['type'] == 'image':
                img = item['image']
                if isinstance(img, np.ndarray):
                    img = Image.fromarray(img).convert("RGB")
                elif isinstance(img, str):
                    if os.path.exists(img):
                        img = Image.open(img).convert("RGB")
                    else:
                        continue
                pil_images.append(img)
                clean_content.append({"type": "image", "image": img})

        if self.use_vllm:
            return self._generate_vllm(clean_content, max_tokens)
        elif "qwen" in self.model_type:
            return self._generate_qwen_transformers(clean_content, max_tokens)
        elif "llava" in self.model_type:
            return self._generate_llava_transformers(clean_content, pil_images, max_tokens)
        
        return "Model backend not supported."

    def _generate_vllm(self, content, max_tokens):
        vllm_messages = []   
        for item in content:
            if item["type"] == "text":
                vllm_messages.append(item)
            elif item["type"] == "image":
                base64_str = image_to_base64(item["image"])
                vllm_messages.append({
                    "type": "image_url",
                    "image_url": {
                        "url": f"data:image/png;base64,{base64_str}"
                    }
                })

        if max_tokens != self.sampling_params.max_tokens:
            self.sampling_params.max_tokens = max_tokens
        
        messages = [{"role": "user", "content": vllm_messages}]
        outputs = self.model.chat(
            messages=messages,
            sampling_params=self.sampling_params,
            use_tqdm=False
        )
        return outputs[0].outputs[0].text

    def _generate_qwen_transformers(self, content, max_tokens):
        messages = [{"role": "user", "content": content}]

        text = self.processor.apply_chat_template(messages, tokenize=False, add_generation_prompt=True)
        image_inputs = []
        for item in content:
            if item["type"] == "image":
                image_inputs.append(item["image"])
        
        if not image_inputs:
            image_inputs = None

        inputs = self.processor(
            text=[text],
            images=image_inputs,
            padding=True,
            return_tensors="pt"
        ).to(self.device)

        with torch.no_grad():
            generated_ids = self.model.generate(
                **inputs, 
                max_new_tokens=max_tokens,
                temperature=0,
                do_sample=False 
            )

        generated_ids_trimmed = [
            out_ids[len(in_ids):] for in_ids, out_ids in zip(inputs.input_ids, generated_ids)
        ]
        output_text = self.processor.batch_decode(
            generated_ids_trimmed, skip_special_tokens=True, clean_up_tokenization_spaces=False
        )
        return output_text[0]

    def _generate_llava_transformers(self, content, pil_images, max_tokens):
        llava_content = []
        for item in content:
            if item['type'] == 'text':
                llava_content.append(item)
            elif item['type'] == 'image':
                llava_content.append({"type": "image"}) 
        
        conversation = [{"role": "user", "content": llava_content}]
        text_prompt = self.processor.apply_chat_template(conversation, add_generation_prompt=True)
        
        inputs = self.processor(
            text=text_prompt, 
            images=pil_images if pil_images else None, 
            return_tensors='pt'
        ).to(self.device, torch.float16)
        
        with torch.no_grad():
            output = self.model.generate(
                **inputs, 
                max_new_tokens=max_tokens,
                temperature=0,
                do_sample=False 
            )
        decoded = self.processor.decode(output[0], skip_special_tokens=True)
        if "assistant\n" in decoded:
            return decoded.split("assistant\n")[-1]
        return decoded

# ==========================================
# 4. Global System Controller
# ==========================================
class GlobalSystem:
    def __init__(self):
        self.vlm = None
        self.localizer = None
        self.rag_agent = None
        self.sam_engine = None
        self.args = None
        self.data_manager = None
        self.is_initialized = False
        self.current_dataset = None
        self.current_ckpt = None
        self.current_localizer = None
        self.current_k_shot = None
        self.current_runtime_signature = None
        self.current_weight_config = None
        self.current_gpu_id = None
        self.current_device = None
        self.weld_runtime = WeldFrontendRuntime()

    def _release_standard_components(self):
        for attr in ("vlm", "localizer", "sam_engine"):
            setattr(self, attr, None)
        self.rag_agent = None
        self.args = None
        self.data_manager = None
        import gc
        gc.collect()
        if torch.cuda.is_available():
            torch.cuda.empty_cache()

    def initialize(
        self,
        model_path,
        ckpt_path,
        graph_root,
        sam_path,
        dataset_root,
        qa_root,
        gpu_id,
        model_type,
        use_vllm,
        gpu_util,
        dataset_name,
        k_shot=1,
        localizer_choice="Auto",
    ):
        try:
            weight_config = _runtime_weight_config(dataset_name, localizer_choice, k_shot, ckpt_path)
            dataset_name = weight_config["dataset_name"]
            k_shot = weight_config["k_shot"]
            localizer_name = weight_config["localizer_name"]
            resolved_ckpt_path = weight_config["checkpoint_path"]
            resolved_save_path = weight_config["save_path"]
            localizer_reload_key = weight_config["signature"]
            if dataset_name == WELD_DATASET:
                self._release_standard_components()
                self.is_initialized = False
                status = self.weld_runtime.initialize(gpu_id, resolved_ckpt_path, sam_path)
                self.is_initialized = True
                self.current_dataset = WELD_DATASET
                self.current_ckpt = localizer_reload_key
                self.current_localizer = "adaptclip"
                self.current_k_shot = 1
                self.current_runtime_signature = localizer_reload_key
                self.current_weight_config = weight_config
                self.current_gpu_id = _normalize_gpu_id(gpu_id)
                self.current_device = self.weld_runtime.device
                return status
            if self.current_dataset == WELD_DATASET:
                self.weld_runtime.clear()
                self.is_initialized = False
                self.current_dataset = None
                self.current_runtime_signature = None
            gpu_index = _normalize_gpu_id(gpu_id)
            if torch.cuda.is_available():
                cuda_count = torch.cuda.device_count()
                if gpu_index < 0 or gpu_index >= cuda_count:
                    raise ValueError(f"GPU {gpu_index} is not visible; PyTorch sees {cuda_count} CUDA device(s).")
                torch.cuda.set_device(gpu_index)
                device = f"cuda:{gpu_index}"
            else:
                device = "cpu"
            gpu_changed = self.current_gpu_id is not None and self.current_gpu_id != gpu_index
            
            print(
                "System Config: "
                f"physical_gpu={gpu_index} | logical_device={device} | dataset={dataset_name} | "
                f"localizer={localizer_name} | k_shot={k_shot}"
            )

            if gpu_changed:
                print(f"GPU changed from {self.current_gpu_id} to {gpu_index}; reloading GPU-bound runtime components.")
                for attr in ("vlm", "localizer", "sam_engine"):
                    component = getattr(self, attr)
                    if component is not None:
                        del component
                        setattr(self, attr, None)
                self.args = None
                self.is_initialized = False
                import gc
                gc.collect()
                if torch.cuda.is_available():
                    torch.cuda.empty_cache()

            # Initialize DatasetManager with the raw path; it handles resolution internally
            self.data_manager = DatasetManager(dataset_root, qa_root, dataset_name)
            
            # Re-init VLM if params changed
            if gpu_changed or self.vlm is None or self.vlm.model_type != model_type or self.vlm.model_path != model_path or self.vlm.use_vllm != use_vllm:
                if self.vlm: 
                    print("Reloading VLM and freeing memory.")
                    del self.vlm
                    import gc
                    gc.collect()
                    if torch.cuda.is_available():
                        torch.cuda.empty_cache()
                
                self.vlm = UnifiedVLMInference(
                    model_path, 
                    device, 
                    model_type, 
                    use_vllm=use_vllm, 
                    gpu_util=gpu_util
                )
            
            # Check if Localizer needs reload
            localizer_needs_reload = (
                self.localizer is None or 
                gpu_changed or
                self.current_dataset != dataset_name or 
                self.current_localizer != localizer_name or
                self.current_ckpt != localizer_reload_key
            )

            if localizer_needs_reload:
                print(f"(Re)Loading {localizer_name} Localizer for dataset: {dataset_name}...")
                if self.localizer:
                    del self.localizer
                    import gc
                    gc.collect()
                    if torch.cuda.is_available():
                        torch.cuda.empty_cache()

                class Args:
                    pass

                self.args = Args()
                self.args.dataset = dataset_name 
                self.args.dataset_root = self.data_manager._resolve_dataset_root()
                self.args.image_size = int(weight_config["image_size"])
                self.args.checkpoint_path = resolved_ckpt_path
                self.args.save_path = resolved_save_path
                self.args.k_shot = k_shot
                self.args.localizer = localizer_name
                
                # Default params
                self.args.features_list = [6, 12, 18, 24]
                self.args.num_visual_finetune_layers = 12
                self.args.depth = 7
                self.args.n_ctx = 11
                self.args.spe = 4
                self.args.w0 = 0.15
                self.args.w1 = 0.35
                self.args.w2 = 0.35
                self.args.w3 = 0.15
                
                if localizer_name == "abound":
                    self.localizer = ABounD_Localizer(self.args, device=device)
                else:
                    self.localizer = AdaptCLIP_Localizer(self.args, device=device)
                
                self.current_dataset = dataset_name
                self.current_ckpt = localizer_reload_key
                self.current_localizer = localizer_name
                self.current_k_shot = k_shot
                self.current_runtime_signature = localizer_reload_key
                self.current_weight_config = weight_config

            # AdaptCLIP can switch shots without a reload. ABounD is selected only
            # for mvtec/visa 1-shot, so a shot change has already reloaded above.
            if self.args is not None:
                self.args.k_shot = k_shot
            if self.localizer is not None:
                self.localizer.k_shot = k_shot
            self.current_k_shot = k_shot
            self.current_runtime_signature = localizer_reload_key
            self.current_weight_config = weight_config
            self.current_gpu_id = gpu_index
            self.current_device = device

            if not self.rag_agent:
                self.rag_agent = SimInspecAgent(None, graph_root, k_shot=1)
            
            if not self.sam_engine and os.path.exists(sam_path): 
                print(f"Loading SAM3 to {device}...")
                self.sam_engine = Sam3Engine(sam_path, device=device)
            
            self.is_initialized = True
            vllm_status = f"ON (Util: {gpu_util})" if (use_vllm and self.vlm.use_vllm) else "OFF"
            shot_mode = f"{k_shot}-shot ({'zero-shot' if k_shot <= 0 else 'few-shot'})"
            localizer_label = "ABounD" if localizer_name == "abound" else "AdaptCLIP"
            weight_note = os.path.basename(weight_config["dataset_weight_paths"].get("memory_bank", resolved_ckpt_path))
            return (
                f"System ready | GPU:{gpu_index} | Model:{model_type} | vLLM:{vllm_status} | "
                f"Dataset:{dataset_name} | Localizer:{localizer_label} {shot_mode} | Weights:{weight_note} | QA:curated"
            )
        except Exception as e:
            if str(dataset_name or "").lower() == WELD_DATASET:
                self.weld_runtime.clear()
                self.is_initialized = False
            import traceback
            traceback.print_exc()
            return f"Init failed: {str(e)}"

    def runtime_match_status(self, dataset_name, localizer_choice, k_shot, ckpt_path, gpu_id=None):
        try:
            expected = _runtime_weight_config(dataset_name, localizer_choice, k_shot, ckpt_path)
        except Exception as e:
            return False, None, str(e)
        if not self.is_initialized:
            return False, expected, "Runtime is not initialized for the selected method configuration."
        if gpu_id is not None:
            try:
                expected_gpu = _normalize_gpu_id(gpu_id)
            except Exception as e:
                return False, expected, f"Invalid GPU selection: {e}"
            if self.current_gpu_id != expected_gpu:
                loaded_gpu = "none" if self.current_gpu_id is None else str(self.current_gpu_id)
                return (
                    False,
                    expected,
                    f"Runtime is loaded on GPU {loaded_gpu}, but the selected GPU is {expected_gpu}. "
                    "Click Initialize runtime before running GLLS.",
                )
        if self.current_runtime_signature != expected["signature"]:
            loaded = _runtime_weight_summary(self.current_weight_config) if self.current_weight_config else "No runtime loaded."
            expected_text = _runtime_weight_summary(expected)
            return (
                False,
                expected,
                "Runtime weights do not match the selected dataset/localizer/shot.\n\n"
                f"Loaded runtime:\n{loaded}\n\n"
                f"Selected runtime:\n{expected_text}\n\n"
                "Click Initialize runtime before running GLLS.",
            )
        return True, expected, "Runtime matches the selected method configuration."

global_sys = GlobalSystem()

class StreamableMCTS(MCTSQuestionSample):
    def __init__(self, row, args, inference_engine, localizer, rag_agent, rag_blocks, sam_engine, update_callback):
        super().__init__(row, args, inference_engine, localizer, rag_agent=rag_agent, rag_blocks=rag_blocks, sam_engine=sam_engine)
        self.update_callback = update_callback 

    async def get_anomaly_heatmap(self, image):
        heatmap, obj_name = await super().get_anomaly_heatmap(image)
        
        if heatmap is None or np.sum(heatmap) == 0:
             await self.update_callback("log", "Warning: localizer returned empty heatmap (check dataset config).")
             return heatmap, obj_name

        try:
            hm_min = heatmap.min()
            hm_max = heatmap.max()
            denominator = hm_max - hm_min + 1e-8
            
            hm_norm = (heatmap - hm_min) / denominator
            hm_uint8 = np.uint8(255 * hm_norm)
            hm_color = cv2.applyColorMap(hm_uint8, cv2.COLORMAP_JET)
            hm_color = cv2.cvtColor(hm_color, cv2.COLOR_BGR2RGB)
            
            img_np = np.array(image.resize((self.image_width, self.image_height)))
            heatmap_resized = cv2.resize(hm_color, (self.image_width, self.image_height))
            
            overlay = cv2.addWeighted(img_np, 0.6, heatmap_resized, 0.4, 0)
            
            await self.update_callback("heatmap", Image.fromarray(overlay))
        except Exception as e:
            print(f"Heatmap Vis Error: {e}")
            
        return heatmap, obj_name

    async def simulation(self, node):
        score = await super().simulation(node)
        if node.state['depth'] > 0:
            await self.update_callback("log", f"MCTS depth {node.state['depth']} searching; score={score:.2f}")
        return score

# ==========================================
# 5. Async Logic Stream
# ==========================================

def _parse_answer_key(response):
    response = _as_text(response)
    match = re.search(r"The correct answer is\s*\(?([A-D])\)?", response, flags=re.IGNORECASE)
    if match:
        return match.group(1).upper()
    cands = re.findall(r"\b([A-D])\b", response)
    return cands[-1].upper() if cands else ""

def _compact_json(value, max_chars=900):
    try:
        text = json.dumps(value, ensure_ascii=False, sort_keys=True)
    except TypeError:
        text = _as_text(value)
    if len(text) > max_chars:
        return text[:max_chars].rstrip() + "\n..."
    return text

def _short_text(value, max_chars=180):
    text = _as_text(value).strip()
    if len(text) > max_chars:
        return text[:max_chars].rstrip() + "..."
    return text

def _source_label(path_value):
    text = _as_text(path_value)
    if not text:
        return ""
    normalized = text.replace("\\", "/")
    for marker in ("/graph_index/", "/text_knowledge/"):
        if marker in normalized:
            return normalized.split(marker, 1)[1]
    return os.path.basename(normalized) or normalized

def _method_trace_markdown(debug_meta):
    debug_meta = debug_meta or {}
    participation = summarize_method_participation(debug_meta)
    rag_summary = participation.get("rag_source_backed_summary", {})
    mcts_budget = debug_meta.get("mcts_budget_config", {}) or {}
    mcts_search = debug_meta.get("mcts_search_summary", {}) or {}
    mcts_action_trace = debug_meta.get("mcts_action_trace", []) or []
    sam_attempts = debug_meta.get("sam_text_prompt_attempts", []) or []
    sam_hits = debug_meta.get("sam_text_prompt_hits", []) or []
    sam_scores = debug_meta.get("sam_mask_scores", []) or []
    rag_blocks = debug_meta.get("rag_block_provenance", []) or []
    graph_paths = []
    text_paths = []
    for block in rag_blocks:
        if not isinstance(block, dict):
            continue
        graph_path = block.get("graph_cache_path")
        source_path = block.get("source_json_path")
        if graph_path and graph_path not in graph_paths:
            graph_paths.append(graph_path)
        if source_path and source_path not in text_paths:
            text_paths.append(source_path)

    lines = [
        "#### Method Trace",
        f"- MCTS: `{participation.get('mcts_participation_status', 'unknown')}` | "
        f"budget `{_compact_json(mcts_budget, 220)}` | actions `{len(mcts_action_trace)}`",
        f"- Region evidence: proposals `{debug_meta.get('region_proposal_count', 0)}` | "
        f"crops `{debug_meta.get('crop_count', 0)}` | prompt-visible `{debug_meta.get('prompt_visible_crop_count', debug_meta.get('crop_count', 0))}`",
        f"- Heatmap: peak `{debug_meta.get('heatmap_peak_score', 0)}` | threshold `{debug_meta.get('anomaly_threshold', 0)}` | source `{debug_meta.get('threshold_source', 'unknown')}`",
        f"- SAM3: `{participation.get('sam_participation_status', 'unknown')}` | "
        f"text attempts `{len(sam_attempts)}` | hits `{len(sam_hits)}` | mask audits `{len(sam_scores)}`",
        f"- PVLA/RAG: blocks `{rag_summary.get('block_count', 0)}` | graph sources `{rag_summary.get('graph_cache_path_count', 0)}` | "
        f"text JSON sources `{rag_summary.get('source_json_path_count', 0)}` | visual cutouts `{rag_summary.get('visual_reference_source_backed_count', 0)}`",
    ]
    if mcts_search:
        lines.append(f"- Search summary: `{_compact_json(mcts_search, 280)}`")
    if graph_paths:
        lines.append(f"- Graph cache: `{_source_label(graph_paths[0])}`" + (f" plus {len(graph_paths) - 1} more" if len(graph_paths) > 1 else ""))
    if text_paths:
        lines.append(f"- Text knowledge: `{_source_label(text_paths[0])}`" + (f" plus {len(text_paths) - 1} more" if len(text_paths) > 1 else ""))
    return "\n".join(lines)

def _first_items(values, limit=4):
    items = []
    for value in values or []:
        if value not in items:
            items.append(value)
        if len(items) >= limit:
            break
    return items

def _html(value):
    return html.escape(_as_text(value), quote=True)

def _display_status(value):
    return _as_text(value).replace("_", " ")

def _top_count_text(counts, limit=4):
    if not isinstance(counts, dict) or not counts:
        return "none"
    pairs = sorted(counts.items(), key=lambda item: item[1], reverse=True)[:limit]
    return ", ".join(f"{key} x{value}" for key, value in pairs)

def _metric_html(label, value, tone=""):
    tone_class = f" {tone}" if tone else ""
    return (
        f'<div class="method-kpi{tone_class}">'
        f"<span>{_html(label)}</span><strong>{_html(value)}</strong>"
        "</div>"
    )

def _item_list_html(items, empty_text):
    rows = [f"<li>{item}</li>" for item in items if item]
    if not rows:
        rows = [f"<li>{_html(empty_text)}</li>"]
    return "<ul>" + "".join(rows) + "</ul>"

def _chip_html(label, value, tone=""):
    tone_class = f" {tone}" if tone else ""
    return (
        f'<span class="evidence-chip{tone_class}">'
        f"<b>{_html(label)}</b><em>{_html(value)}</em>"
        "</span>"
    )

def _mini_list_html(items, empty_text, limit=5):
    rows = []
    for item in _first_items(items, limit):
        rows.append(f"<li>{_html(item)}</li>")
    if not rows:
        rows = [f"<li>{_html(empty_text)}</li>"]
    return "<ul>" + "".join(rows) + "</ul>"

def _load_pickle_on_cpu(path):
    class _CPUUnpickler(pickle.Unpickler):
        def find_class(self, module, name):
            if module == "torch.storage" and name == "_load_from_bytes":
                return lambda payload: torch.load(
                    BytesIO(payload),
                    map_location="cpu",
                    weights_only=False,
                )
            return super().find_class(module, name)

    with open(path, "rb") as handle:
        return _CPUUnpickler(handle).load()


def load_pvla_graph_summary(category, graph_path=None):
    category = _as_text(category).strip()
    summary = {
        "category": category or "selected category",
        "path": "",
        "status": "waiting for a loaded category",
        "node_count": 0,
        "edge_count": 0,
        "nodes": [],
        "edges": [],
        "node_details": [],
        "edge_details": [],
        "source": "",
        "dataset": "",
    }
    if not category:
        return summary

    graph_path = str(graph_path or os.path.join(DEFAULT_GRAPH_ROOT, f"{category}_graph.pkl"))
    summary["path"] = graph_path
    if not os.path.exists(graph_path):
        summary["status"] = "graph cache not found"
        return summary

    try:
        payload = _load_pickle_on_cpu(graph_path)
        graph = payload.get("graph") if isinstance(payload, dict) else payload
        metadata = payload.get("source_metadata", {}) if isinstance(payload, dict) else {}
        node_names = payload.get("node_names", []) if isinstance(payload, dict) else []

        if hasattr(graph, "nodes"):
            graph_nodes = list(graph.nodes(data=True))
            graph_edges = list(graph.edges(data=True))
            names = node_names or [str(node) for node, _ in graph_nodes]
            node_details = []
            for node_id, data in graph_nodes:
                data = data or {}
                label = data.get("label") or data.get("short_name") or _source_label(str(node_id)) or str(node_id)
                image_paths = data.get("image_paths") or data.get("images") or []
                if isinstance(image_paths, str):
                    image_paths = [image_paths]
                node_details.append(
                    {
                        "id": str(node_id),
                        "label": str(label),
                        "type": str(data.get("type") or "node"),
                        "images": [str(path) for path in image_paths[:4]],
                    }
                )
            edge_details = []
            for src, dst, data in graph_edges:
                data = data or {}
                edge_details.append(
                    {
                        "source": str(src),
                        "target": str(dst),
                        "relation": str(data.get("relation") or "linked"),
                    }
                )
            summary["node_details"] = node_details
            summary["edge_details"] = edge_details
            summary["nodes"] = [str(name) for name in names[:10]]
            summary["edges"] = [
                f"{_source_label(str(src))} -> {_source_label(str(dst))}"
                for src, dst, *_ in graph_edges[:10]
            ]
            summary["node_count"] = len(graph_nodes)
            summary["edge_count"] = len(graph_edges)
        else:
            summary["nodes"] = [str(node) for node in node_names[:10]]
            summary["node_count"] = len(node_names)

        summary["source"] = _source_label(metadata.get("source_json_path", ""))
        summary["dataset"] = _as_text(metadata.get("dataset", ""))
        summary["status"] = "loaded"
    except Exception as e:
        summary["status"] = f"graph read failed: {_short_text(e, 90)}"
    return summary

def _atlas_node_tone(node_type):
    text = _as_text(node_type).lower()
    if "root" in text or "class" in text:
        return "root"
    if "defect" in text or "anomaly" in text:
        return "defect"
    if "normal" in text:
        return "source"
    if "region" in text or "part" in text:
        return "region"
    if "image" in text or "source" in text:
        return "source"
    return "node"


_ATLAS_TONE_KIND = {
    "root": "class",
    "region": "part",
    "source": "normal",
    "defect": "defect",
    "node": "node",
}


def _atlas_clean_label(node, fallback="node"):
    """Region nodes carry an empty label, so fall back to a tidied id."""
    raw = _as_text(node.get("label")).strip()
    if not raw:
        raw = _as_text(node.get("id")).strip()
    raw = raw.replace("_", " ").strip()
    if raw.lower().startswith("part "):
        raw = raw[5:].strip() or raw
    return raw or fallback


def _box_edge_point(cx, cy, half_w, half_h, tx, ty):
    """Point where the line from (cx, cy) toward (tx, ty) crosses the node box."""
    dx, dy = tx - cx, ty - cy
    if abs(dx) < 1e-6 and abs(dy) < 1e-6:
        return cx, cy
    scale = float("inf")
    if abs(dx) > 1e-6:
        scale = min(scale, half_w / abs(dx))
    if abs(dy) > 1e-6:
        scale = min(scale, half_h / abs(dy))
    return cx + dx * scale, cy + dy * scale


def _atlas_svg_node(node, cx, cy, width, height):
    """Render a single pill node centered on (cx, cy)."""
    tone = _atlas_node_tone(node.get("type", "node"))
    label = _short_text(_atlas_clean_label(node), 40 if tone == "root" else 26)
    kind = _ATLAS_TONE_KIND.get(tone, "node")
    x = cx - width / 2.0
    y = cy - height / 2.0
    return (
        f'<foreignObject x="{x:.1f}" y="{y:.1f}" width="{width:.1f}" height="{height:.1f}">'
        f'<div xmlns="http://www.w3.org/1999/xhtml" class="atlas-node {tone}" title="{_html(label)}">'
        '<i class="atlas-dot"></i>'
        f'<span class="atlas-node-text"><em>{_html(kind)}</em><strong>{_html(label)}</strong></span>'
        "</div></foreignObject>"
    )


def render_pvla_atlas_graph(atlas):
    """Lay out the PVLA category graph as a hub-and-ring diagram.

    Class node in the center, part/region nodes on an inner ring, and defect
    patterns evenly spaced on an outer ring (ordered by parent part so links
    fan out cleanly). Positions are computed so node boxes never overlap.
    """
    nodes = atlas.get("node_details") or []
    category = atlas.get("category") or "selected category"
    if not nodes:
        nodes = [{"id": category, "label": category, "type": "root", "images": []}]

    # ---- Bucket nodes by semantic role (Class / Part / Defect) ----
    root = next((n for n in nodes if _atlas_node_tone(n.get("type")) == "root"), None)
    if root is None:
        root = nodes[0]
    root_id = root.get("id")

    regions, defects = [], []
    for n in nodes:
        if n.get("id") == root_id:
            continue
        tone = _atlas_node_tone(n.get("type"))
        if tone == "defect":
            defects.append(n)
        elif tone in ("region", "source"):
            regions.append(n)
        # generic / composite (untyped) nodes are skipped to keep the view legible

    MAX_REGIONS, MAX_DEFECTS = 6, 14
    region_total, defect_total = len(regions), len(defects)
    regions = regions[:MAX_REGIONS]
    defects = defects[:MAX_DEFECTS]
    region_ids = {r.get("id") for r in regions}
    defect_ids = {d.get("id") for d in defects}

    # ---- Attach each defect to a parent region (edges first, id-prefix fallback) ----
    edge_details = atlas.get("edge_details") or []
    parent_of = {}
    for edge in edge_details:
        s, d = edge.get("source"), edge.get("target")
        rel = _as_text(edge.get("relation")).lower()
        if "anomaly" not in rel and "region" not in rel and "has" not in rel:
            continue
        if d in defect_ids and s in region_ids:
            parent_of.setdefault(d, s)
        elif s in defect_ids and d in region_ids:
            parent_of.setdefault(s, d)
    for d in defects:
        did = d.get("id")
        if did in parent_of:
            continue
        best = None
        for r in regions:
            rid = _as_text(r.get("id"))
            if rid and _as_text(did).startswith(rid):
                if best is None or len(rid) > len(_as_text(best)):
                    best = r.get("id")
        if best:
            parent_of[did] = best

    region_index = {r.get("id"): i for i, r in enumerate(regions)}
    defects_ordered = sorted(
        defects,
        key=lambda d: (region_index.get(parent_of.get(d.get("id")), len(regions)),),
    )

    # ---- Geometry (viewBox units) ----
    W, H = 900.0, 620.0
    cx, cy = W / 2.0, H / 2.0
    R_REGION, R_DEFECT = 152.0, 288.0
    ROOT_W, ROOT_H = 150.0, 52.0
    REGION_W, REGION_H = 138.0, 44.0
    DEFECT_W, DEFECT_H = 118.0, 40.0
    TOP = -math.pi / 2.0

    pos, half = {}, {}
    pos[root_id] = (cx, cy)
    half[root_id] = (ROOT_W / 2.0, ROOT_H / 2.0)

    # Defects: even angular spread on the outer ring (grouped order keeps links short)
    defect_angle = {}
    n_def = max(1, len(defects_ordered))
    for i, d in enumerate(defects_ordered):
        ang = TOP + 2.0 * math.pi * i / n_def
        defect_angle[d.get("id")] = ang
        pos[d.get("id")] = (cx + R_DEFECT * math.cos(ang), cy + R_DEFECT * math.sin(ang))
        half[d.get("id")] = (DEFECT_W / 2.0, DEFECT_H / 2.0)

    # Regions: placed at the angular centroid of their defects (fallback: even spread)
    fallback_slots = [TOP + 2.0 * math.pi * k / max(1, len(regions)) for k in range(len(regions))]
    for i, r in enumerate(regions):
        rid = r.get("id")
        child_angles = [defect_angle[d.get("id")] for d in defects_ordered if parent_of.get(d.get("id")) == rid]
        if child_angles:
            sxc = sum(math.cos(a) for a in child_angles)
            syc = sum(math.sin(a) for a in child_angles)
            ang = math.atan2(syc, sxc) if (abs(sxc) > 1e-9 or abs(syc) > 1e-9) else fallback_slots[i]
        else:
            ang = fallback_slots[i]
        pos[rid] = (cx + R_REGION * math.cos(ang), cy + R_REGION * math.sin(ang))
        half[rid] = (REGION_W / 2.0, REGION_H / 2.0)

    displayed = set(pos.keys())

    # ---- Edges (hierarchy vs distinct_from) ----
    hierarchy, diffs = [], []
    have_edges = False
    for edge in edge_details:
        s, d = edge.get("source"), edge.get("target")
        if s not in displayed or d not in displayed or s == d:
            continue
        rel = _as_text(edge.get("relation")).lower()
        if "distinct" in rel or "diff" in rel:
            diffs.append((s, d))
        else:
            hierarchy.append((s, d))
        have_edges = True
    if not have_edges:
        hierarchy = [(root_id, r.get("id")) for r in regions]
        for d in defects_ordered:
            hierarchy.append((parent_of.get(d.get("id")) or root_id, d.get("id")))
    existing = {frozenset((s, d)) for s, d in hierarchy}
    for r in regions:
        if frozenset((root_id, r.get("id"))) not in existing:
            hierarchy.append((root_id, r.get("id")))

    def _link(s, d, cls):
        sx, sy = pos[s]
        dx, dy = pos[d]
        shw, shh = half[s]
        dhw, dhh = half[d]
        x1, y1 = _box_edge_point(sx, sy, shw + 2, shh + 2, dx, dy)
        x2, y2 = _box_edge_point(dx, dy, dhw + 5, dhh + 5, sx, sy)
        return f'<line class="atlas-link {cls}" x1="{x1:.1f}" y1="{y1:.1f}" x2="{x2:.1f}" y2="{y2:.1f}" />'

    link_markup = [_link(s, d, "hier") for s, d in hierarchy[:48]]
    diff_markup = [_link(s, d, "diff") for s, d in diffs[:8]]

    rings = (
        f'<circle class="atlas-ring" cx="{cx:.0f}" cy="{cy:.0f}" r="{R_REGION:.0f}" />'
        f'<circle class="atlas-ring" cx="{cx:.0f}" cy="{cy:.0f}" r="{R_DEFECT:.0f}" />'
    )

    node_markup = [_atlas_svg_node(root, cx, cy, ROOT_W, ROOT_H)]
    for r in regions:
        x, y = pos[r.get("id")]
        node_markup.append(_atlas_svg_node(r, x, y, REGION_W, REGION_H))
    for d in defects_ordered:
        x, y = pos[d.get("id")]
        node_markup.append(_atlas_svg_node(d, x, y, DEFECT_W, DEFECT_H))

    legend = (
        '<div class="atlas-legend">'
        '<span class="atlas-leg root"><i></i>Class</span>'
        '<span class="atlas-leg region"><i></i>Part / region</span>'
        '<span class="atlas-leg defect"><i></i>Defect pattern</span>'
        '<span class="atlas-leg diff"><i></i>distinct_from</span>'
        "</div>"
    )

    overflow_bits = []
    if region_total > len(regions):
        overflow_bits.append(f"+{region_total - len(regions)} more parts")
    if defect_total > len(defects):
        overflow_bits.append(f"+{defect_total - len(defects)} more defects")
    overflow = (
        '<p class="atlas-overflow">'
        f"{len(regions)} parts · {len(defects)} defect patterns"
        + (" · " + ", ".join(overflow_bits) if overflow_bits else "")
        + "</p>"
    )

    return (
        '<div class="atlas-graph-wrap" aria-label="PVLA graph overview">'
        + legend
        + f'<svg class="atlas-graph-svg" viewBox="0 0 {int(W)} {int(H)}" role="img" preserveAspectRatio="xMidYMid meet">'
        + '<defs><marker id="atlas-arrow" markerWidth="9" markerHeight="9" refX="8" refY="4" '
        'orient="auto" markerUnits="userSpaceOnUse">'
        + '<path d="M0,0 L9,4 L0,8 Z"></path></marker></defs>'
        + rings
        + "".join(link_markup)
        + "".join(diff_markup)
        + "".join(node_markup)
        + "</svg>"
        + overflow
        + "</div>"
    )

def _resolve_pvla_image_path(stored_path, graph_path=""):
    """Re-root a build-machine image path (e.g. /home/<other>/.../img/...) onto
    the local img/ directory that sits next to the loaded graph cache."""
    p = _as_text(stored_path).strip()
    if not p:
        return None
    if os.path.exists(p):
        return p
    norm = p.replace("\\", "/")
    suffix = None
    for marker in ("/databases/img/", "/img/"):
        if marker in norm:
            suffix = norm.split(marker, 1)[1]
            break
    if not suffix:
        return None
    bases = []
    gp = _as_text(graph_path).replace("\\", "/")
    if gp:
        bases.append(os.path.join(os.path.dirname(os.path.dirname(gp)), "img"))
        bases.append(os.path.join(os.path.dirname(gp), "img"))
    bases.append(os.path.join(os.path.dirname(_as_text(DEFAULT_GRAPH_ROOT)), "img"))
    for base in bases:
        candidate = os.path.join(base, suffix)
        if os.path.exists(candidate):
            return candidate
    return None


def _pvla_thumb_data_uri(path, max_px=150):
    """Load, downscale, and base64-encode an image so it renders inside gr.HTML."""
    try:
        with Image.open(path) as im:
            im = im.convert("RGB")
            im.thumbnail((max_px, max_px))
            bio = BytesIO()
            im.save(bio, format="JPEG", quality=80)
        return "data:image/jpeg;base64," + base64.b64encode(bio.getvalue()).decode("utf-8")
    except Exception:
        return None


def _atlas_part_crops_html(atlas):
    """Show the actual reference crops linked to part/region nodes. Defect-pattern
    nodes are text-only by design, so they are summarized as a count, not listed."""
    graph_path = atlas.get("path", "")
    cards = []
    defect_count = 0
    for node in atlas.get("node_details") or []:
        tone = _atlas_node_tone(node.get("type"))
        if tone == "defect":
            defect_count += 1
            continue
        if tone not in ("region", "source"):
            continue
        uri = None
        for raw in node.get("images", []):
            local = _resolve_pvla_image_path(raw, graph_path)
            if local:
                uri = _pvla_thumb_data_uri(local)
                if uri:
                    break
        label = _atlas_clean_label(node)
        if uri:
            inner = f'<img src="{uri}" alt="{_html(label)}" loading="lazy" />'
        else:
            inner = '<span class="atlas-crop-missing">crop unavailable</span>'
        cards.append(
            f'<figure class="atlas-crop"><div class="atlas-crop-img">{inner}</div>'
            f"<figcaption>{_html(label)}</figcaption></figure>"
        )
    if not cards:
        return '<p class="atlas-crop-empty">Part reference crops appear once a category graph is loaded.</p>'
    note = ""
    if defect_count:
        note = (
            f'<p class="atlas-crop-note">{defect_count} defect patterns are text-only '
            "(visual signature &amp; contrast); they have no reference crop by design.</p>"
        )
    return '<div class="atlas-crop-grid">' + "".join(cards) + "</div>" + note


def _atlas_detail_html(atlas, rag_blocks, regions, source, dataset, graph_sources):
    edge_rows = []
    for edge in (atlas.get("edge_details") or [])[:12]:
        edge_rows.append(
            f"<li><code>{_html(_source_label(edge.get('source')))}</code> "
            f"-> <code>{_html(_source_label(edge.get('target')))}</code> "
            f"<em>{_html(edge.get('relation'))}</em></li>"
        )
    if not edge_rows:
        edge_rows.append("<li>No graph edges loaded yet.</li>")

    prov_rows = []
    regions_text = ", ".join(_first_items(regions, 5))
    graph_cache = ", ".join(_first_items(graph_sources, 2)) or _source_label(atlas.get("path", ""))
    for label, value in [
        ("Knowledge source", source),
        ("Dataset", dataset),
        ("RAG regions", regions_text),
        ("Graph cache", graph_cache),
    ]:
        value = _as_text(value).strip()
        if value:
            prov_rows.append(f"<p><b>{_html(label)}</b><br>{_html(value)}</p>")
    if not prov_rows:
        prov_rows.append(f'<p><b>Status</b><br>{_html(atlas.get("status", "not loaded"))}</p>')

    return (
        '<details class="atlas-details">'
        "<summary>Open PVLA part crops, edges, and recall provenance</summary>"
        '<div class="atlas-detail-grid">'
        "<div><h4>Part reference crops</h4>" + _atlas_part_crops_html(atlas) + "</div>"
        '<div><h4>Graph edges</h4><ul>' + "".join(edge_rows) + "</ul></div>"
        "<div><h4>Recall provenance</h4>" + "".join(prov_rows) + "</div>"
        "</div></details>"
    )

def _atlas_panel_html(category, rag_blocks=None, graph_path=None):
    atlas = load_pvla_graph_summary(category, graph_path=graph_path)
    rag_blocks = rag_blocks or []
    regions = []
    graph_sources = []
    text_sources = []
    for block in rag_blocks:
        if not isinstance(block, dict):
            continue
        if block.get("region"):
            regions.append(str(block["region"]))
        if block.get("graph_cache_path"):
            graph_sources.append(_source_label(block["graph_cache_path"]))
        if block.get("source_json_path"):
            text_sources.append(_source_label(block["source_json_path"]))

    source = atlas.get("source") or ", ".join(_first_items(text_sources, 2))
    dataset = atlas.get("dataset")
    graph_bits = (
        _chip_html("category", atlas["category"], "pvla")
        + _chip_html("nodes", atlas.get("node_count", 0), "pvla")
        + _chip_html("edges", atlas.get("edge_count", 0), "pvla")
        + _chip_html("status", atlas.get("status", "unknown"), "pvla")
    )
    return (
        '<section class="stream-panel atlas-panel">'
        '<div class="stream-heading"><div><p class="eyebrow">offline atlas</p>'
        '<h3>PVLA graph-shaped knowledge</h3>'
        '<p>Part, normal-reference, and defect knowledge are recalled as a category graph before the final verifier sees the case.</p>'
        '</div></div>'
        f'<div class="chip-row">{graph_bits}</div>'
        + render_pvla_atlas_graph(atlas)
        + _atlas_detail_html(atlas, rag_blocks, regions, source, dataset, graph_sources)
        + '</section>'
    )

def render_dual_stream_placeholder(category=None):
    return (
        '<div class="dual-board-html">'
        '<div class="board-head"><div><p class="eyebrow">demo board</p>'
        '<h2>Dual-stream verification is ready for a QA sample</h2>'
        '<p>Select a dataset/category on the left, then run GLLS to fill the board with graph knowledge, heatmap search, local crops, and the final verifier decision.</p>'
        '</div><span class="process-badge">Waiting for run</span></div>'
        '<div class="board-grid">'
        + _atlas_panel_html(category)
        + '<section class="stream-panel global-panel"><div class="stream-heading"><p class="eyebrow">stream 1</p><h3>Global & logic stream</h3><p>Whole-image reasoning and structural checks appear here after runtime execution.</p></div></section>'
        + '<section class="stream-panel action-panel"><div class="stream-heading"><p class="eyebrow">stream 2</p><h3>Fine-grained & actions stream</h3><p>AdaptCLIP heatmap, MCTS search, and Top-K crops are paired here instead of split across separate cards.</p></div></section>'
        + '<section class="stream-panel fusion-panel"><div class="stream-heading"><p class="eyebrow">fusion</p><h3>Cross-stream evidence handoff</h3><p>Only selected global, local, and PVLA evidence is passed to the Phase-2 verifier.</p></div></section>'
        + '<section class="stream-panel verdict-panel"><div class="stream-heading"><p class="eyebrow">verdict</p><h3>Final answer</h3><p>The constrained multiple-choice result and evidence badges appear here after a run.</p></div></section>'
        '</div></div>'
    )


def _weld_reference_panel_html(result=None):
    result = result or {}
    try:
        manifest = global_sys.weld_runtime.manifest
        counts = manifest.get("counts") or {}
        references = global_sys.weld_runtime.normal_reference_paths()
        graph_path = global_sys.weld_runtime.root / "assets" / "graph" / "gear_weld_graph.pkl"
    except Exception:
        counts = {}
        references = []
        graph_path = ""
    pvla = result.get("pvla") or {}
    reference_count = int(pvla.get("reference_count", 0) or 0)
    similarity = pvla.get("normal_similarity")
    hotspot_xy = pvla.get("hotspot_xy") or []
    best_reference = (pvla.get("best_reference") or {}).get("source_path")
    reference_card = '<p class="atlas-crop-empty">Build weld assets to show the trusted normal reference.</p>'
    displayed_reference = best_reference or (references[-1] if references else "")
    if displayed_reference:
        uri = _pvla_thumb_data_uri(displayed_reference, max_px=260)
        if uri:
            reference_card = (
                '<figure class="weld-reference-card">'
                f'<img src="{uri}" alt="Trusted normal gear-weld reference" />'
                '<figcaption>Retrieved from PDF page 4 trusted-normal atlas</figcaption></figure>'
            )
    return (
        _atlas_panel_html(WELD_CATEGORY, graph_path=graph_path)
        + '<section class="stream-panel pvla-recall-panel">'
        '<div class="stream-heading"><div><p class="eyebrow">online PVLA recall</p>'
        '<h3>Graph recall → trusted-normal visual comparison</h3>'
        '<p>The offline region/defect graph selects the relevant normal atlas before the online hotspot is compared.</p>'
        '</div></div><div class="chip-row">'
        + _chip_html("normal shots", 1, "pvla")
        + _chip_html("PVLA patches", reference_count or "after init", "pvla")
        + _chip_html("review images", int(counts.get("annotated_images", 0)) + int(counts.get("raw_unlabeled_images", 0)), "pvla")
        + _chip_html("normal similarity", f"{float(similarity):.4f}" if similarity is not None else "after run", "pvla")
        + _chip_html("hotspot xy", ", ".join(map(str, hotspot_xy)) if hotspot_xy else "after run", "pvla")
        + '</div><div class="paired-grid"><div class="graph-mini"><strong>Evidence handoff</strong><ul>'
        '<li>Offline: PDF-backed regions, defect nodes, distinctions, and graph edges.</li>'
        '<li>Online: tiled hotspot is embedded and compared with retrieved normal patches.</li>'
        '<li>Decision: compact hotspot uses the fixed operating point; PVLA similarity remains explicit supporting evidence.</li>'
        '</ul></div><div>'
        + reference_card
        + '</div></div></section>'
    )


def render_weld_placeholder():
    return (
        '<div class="dual-board-html">'
        '<div class="board-head"><div><p class="eyebrow">one-category method demo</p>'
        '<h2>GLLS gear-weld inspection is ready</h2>'
        '<p>Select one weld review image, initialize AdaptCLIP/PVLA/SAM3, then run the same dual-stream evidence board.</p>'
        '</div><span class="process-badge">Method walkthrough · not a benchmark</span></div>'
        '<div class="board-grid">'
        + _weld_reference_panel_html()
        + '<section class="stream-panel global-panel"><div class="stream-heading"><p class="eyebrow">stream 1</p><h3>Global weld logic</h3><p>SAM3-assisted geometry checks inspect ring or visible-arc continuity before local scoring.</p></div></section>'
        + '<section class="stream-panel action-panel"><div class="stream-heading"><p class="eyebrow">stream 2</p><h3>Tiled local anomaly evidence</h3><p>High-resolution overlapping AdaptCLIP tiles preserve compact weld defects; no MCTS claim is made on this route.</p></div></section>'
        + '<section class="stream-panel fusion-panel"><div class="stream-heading"><p class="eyebrow">fusion</p><h3>Conservative OR gate</h3><p>A global structural trigger or a fixed-threshold compact hotspot is sufficient for an NG review verdict.</p></div></section>'
        + '<section class="stream-panel verdict-panel"><div class="stream-heading"><p class="eyebrow">verdict</p><h3>Prediction / reference comparison</h3><p>dx is an annotated anomaly; ljtq is an allowed station-2 connection context. Unlabelled images remain unscored.</p></div></section>'
        '</div></div>'
    )


def render_weld_running():
    return (
        '<div class="dual-board-html weld-running-board">'
        '<div class="board-head"><div><p class="eyebrow">live method execution</p>'
        '<h2>Running GLLS…</h2>'
        '<p>The selected image is being processed now. High-resolution tiled inference usually takes tens of seconds.</p>'
        '</div><span class="process-badge running-badge"><i></i>Processing</span></div>'
        '<div class="run-progress-grid">'
        '<div class="run-step active"><b>1</b><span><strong>SAM3 structure</strong><em>Segment weld geometry and audit global rules</em></span></div>'
        '<div class="run-step active"><b>2</b><span><strong>AdaptCLIP tiles</strong><em>Scan overlapping 1600 px local windows</em></span></div>'
        '<div class="run-step active"><b>3</b><span><strong>PVLA recall</strong><em>Compare hotspots with graph-retrieved normal patches</em></span></div>'
        '<div class="run-step active"><b>4</b><span><strong>OR fusion</strong><em>Render prediction, reference, overlay, and crops</em></span></div>'
        '</div><p class="run-wait-note">Please keep this tab open. The evidence board will replace this panel automatically when the run completes.</p>'
        '</div>'
    )


def render_weld_trace(result):
    result = result or {}
    global_result = result.get("global") or {}
    trace = result.get("trace") or {}
    local_trace = trace.get("local_stream") or {}
    rules = global_result.get("rules") or []
    rule_items = [
        f"<b>{_html(rule.get('name'))}</b>: <code>{_html(rule.get('status'))}</code> "
        f"<span>({_html(rule.get('criterion'))})</span>"
        for rule in rules
    ]
    compact = float(local_trace.get("compact_hotspot_score", 0.0) or 0.0)
    similarity = float(local_trace.get("pvla_normal_similarity", 0.0) or 0.0)
    orchestra = float(local_trace.get("orchestra_score", result.get("orchestra_score", 0.0)) or 0.0)
    threshold = float(local_trace.get("alert_threshold", result.get("local_threshold", 0.0)) or 0.0)
    verdict = str(result.get("verdict") or "WAIT")
    reference = str(
        trace.get("reference_ground_truth")
        or trace.get("official_ground_truth")
        or (result.get("reference") or {}).get("ground_truth")
        or "UNKNOWN"
    ).upper()
    if reference in {"OK", "NG"}:
        comparison = "MATCH" if verdict == reference else "MISS" if reference == "NG" else "FALSE ALARM"
    else:
        comparison = "UNSCORED"
    sam_overlay_uri = _image_data_uri(global_result.get("overlay"), max_px=1000)
    sam_overlay = (
        '<figure class="weld-evidence-figure"><img src="'
        + sam_overlay_uri
        + '" alt="SAM3 weld structure overlay" /><figcaption>SAM3 mask + geometry prior + logic status</figcaption></figure>'
        if sam_overlay_uri
        else '<p class="atlas-crop-empty">No SAM3 structure overlay was recorded.</p>'
    )
    crop_cards = []
    for index, item in enumerate(result.get("hotspot_gallery") or [], start=1):
        image_like = item[0] if isinstance(item, (list, tuple)) and item else item
        caption = item[1] if isinstance(item, (list, tuple)) and len(item) > 1 else f"Online crop {index}"
        uri = _image_data_uri(image_like, max_px=620)
        if uri:
            crop_cards.append(
                '<figure class="weld-evidence-figure"><img src="'
                + uri
                + f'" alt="Online anomaly crop {index}" /><figcaption>{_html(caption)}</figcaption></figure>'
            )
    online_crops = (
        '<div class="weld-crop-grid">' + "".join(crop_cards) + "</div>"
        if crop_cards
        else '<p class="atlas-crop-empty">Online crops appear after the tiled detector runs.</p>'
    )
    annotation_uri = _image_data_uri(result.get("annotation_view"), max_px=1000)
    annotation_figure = (
        '<figure class="weld-evidence-figure"><img src="'
        + annotation_uri
        + '" alt="Reference annotation overlay" /><figcaption>Reference overlay: red dx anomaly / blue ljtq allowed context</figcaption></figure>'
        if annotation_uri
        else ""
    )
    global_panel = (
        '<section class="stream-panel global-panel"><div class="stream-heading"><div>'
        '<p class="eyebrow">stream 1</p><h3>SAM3 logic anomaly check</h3>'
        '<p>The segmentation mask is constrained by the station geometry prior, then audited by explicit continuity and thickness rules.</p>'
        '</div></div><div class="chip-row">'
        + _chip_html("station", global_result.get("station", "unknown"), "global")
        + _chip_html("mask", global_result.get("mask_source", "unknown"), "global")
        + _chip_html("SAM3 score", f"{float(global_result.get('sam3_score') or 0.0):.4f}", "global")
        + _chip_html("mask quality", f"{float(global_result.get('sam3_quality') or 0.0):.4f}", "global")
        + _chip_html("structure", "reliable" if global_result.get("reliable_structure") else "insufficient", "global")
        + _chip_html("logic", "NG" if result.get("global_logic_ng") else "pass", "global")
        + '</div>'
        + sam_overlay
        + '<details open><summary>Structural rule audit</summary>'
        + _item_list_html(rule_items, "No global rule result was recorded.")
        + '</details></section>'
    )
    local_panel = (
        '<section class="stream-panel action-panel"><div class="stream-heading"><div>'
        '<p class="eyebrow">stream 2</p><h3>Tiled AdaptCLIP + PVLA</h3>'
        '<p>The compact hotspot is the local decision score; PVLA reports graph-retrieved normal similarity alongside it without an uncalibrated suppression multiplier.</p>'
        '</div></div><div class="chip-row">'
        + _chip_html("compact hotspot", f"{compact:.4f}", "action")
        + _chip_html("normal similarity", f"{similarity:.4f}", "pvla")
        + _chip_html("local score", f"{orchestra:.4f}", "action")
        + _chip_html("threshold", f"{threshold:.4f}", "action")
        + '</div>'
        + _item_list_html([
            f"Local trigger: <code>{'NG' if result.get('local_ng') else 'pass'}</code>",
            f"Top tiled regions: <code>{len((result.get('local') or {}).get('top_tiles', []) or [])}</code>",
            "Operating point: <code>fixed compact-hotspot threshold from existing offline validation</code>",
        ], "No local evidence was recorded.")
        + '</section>'
    )
    crop_panel = (
        '<section class="stream-panel online-crop-panel"><div class="stream-heading"><div>'
        '<p class="eyebrow">online local evidence</p><h3>Online anomaly crops</h3>'
        '<p>These crops are cut from the current image at runtime. Each caption records the hotspot center and compact response.</p>'
        '</div></div>'
        + online_crops
        + '</section>'
    )
    fusion_panel = (
        '<section class="stream-panel fusion-panel"><div class="stream-heading"><div>'
        '<p class="eyebrow">fusion</p><h3>Global logic NG OR local AdaptCLIP NG</h3>'
        '<p>The two streams remain inspectable and the fusion rule is explicit.</p>'
        '</div></div><div class="chip-row">'
        + _chip_html("global", "NG" if result.get("global_logic_ng") else "pass", "global")
        + _chip_html("local", "NG" if result.get("local_ng") else "pass", "action")
        + _chip_html("fusion", verdict, "")
        + '</div></section>'
    )
    verdict_panel = (
        '<section class="stream-panel verdict-panel"><div class="stream-heading"><div>'
        '<p class="eyebrow">prediction vs reference</p><h3>Prediction '
        + _html(verdict)
        + ' / Ground truth '
        + _html(reference)
        + ' / '
        + _html(comparison)
        + '</h3><p>Station-2 mapping: <code>dx → NG defect</code>; <code>ljtq → OK allowed connection context</code>. '
        'This reference is backed by the official inspection PDFs and archive polygons, but the archive still has no official benchmark split.</p>'
        '</div></div>'
        + annotation_figure
        + '</section>'
    )
    return (
        '<div class="dual-board-html"><div class="board-head"><div>'
        '<p class="eyebrow">trace-backed weld demo</p><h2>Dual-stream gear-weld inspection</h2>'
        '<p>The board below is rendered from the weld trace produced by this run.</p>'
        '</div><span class="process-badge">Trace-backed · '
        + _html(verdict)
        + '</span></div><div class="board-grid">'
        + _weld_reference_panel_html(result)
        + global_panel
        + local_panel
        + crop_panel
        + fusion_panel
        + verdict_panel
        + '</div></div>'
    )

def _mcts_action_viz_html(mcts_search, action_trace, mcts_status):
    """Render the MCTS action budget as labelled bars (Move / Zoom / Stop)."""
    counts = (mcts_search or {}).get("expanded_action_counts", {}) or {}
    families = {"Move": 0, "Zoom": 0, "Stop": 0}
    for key, value in counts.items():
        name = str(key).lower()
        try:
            value = int(value)
        except (TypeError, ValueError):
            value = 0
        if "move" in name:
            families["Move"] += value
        elif "zoom" in name:
            families["Zoom"] += value
        elif "stop" in name or "terminal" in name:
            families["Stop"] += value
    terminal = int((mcts_search or {}).get("terminal_iterations", 0) or 0)
    if terminal and families["Stop"] == 0:
        families["Stop"] = terminal
    iterations = int((mcts_search or {}).get("iterations_recorded", len(action_trace or [])) or 0)
    max_depth = (mcts_search or {}).get("max_path_depth", (mcts_search or {}).get("max_selection_depth", ""))

    if iterations == 0 and sum(families.values()) == 0 and not counts:
        return (
            '<div class="mcts-viz empty">MCTS search was not triggered for this task '
            f"(<code>{_html(mcts_status)}</code>); the global stream answered directly.</div>"
        )

    icons = {"Move": "&#8596;", "Zoom": "&#8853;", "Stop": "&#9632;"}
    peak = max(1, max(families.values()))
    bars = []
    for name in ("Move", "Zoom", "Stop"):
        count = families[name]
        pct = int(round(100 * count / peak))
        bars.append(
            f'<div class="mcts-act {name.lower()}">'
            f'<span class="mcts-act-h"><i>{icons[name]}</i>{name}<b>{count}</b></span>'
            f'<span class="mcts-bar"><span style="width:{pct}%"></span></span></div>'
        )
    meta = (
        f'<div class="mcts-viz-meta">{iterations} search iterations'
        + (f" &middot; max depth {_html(max_depth)}" if _as_text(max_depth) else "")
        + "</div>"
    )
    return '<div class="mcts-viz">' + "".join(bars) + meta + "</div>"


def render_dual_stream_trace(debug_meta, final_prompt="", category=""):
    debug_meta = debug_meta or {}
    participation = summarize_method_participation(debug_meta)
    mcts_budget = debug_meta.get("mcts_budget_config", {}) or {}
    mcts_search = debug_meta.get("mcts_search_summary", {}) or {}
    crop_audit = debug_meta.get("crop_evidence_audit", []) or []
    sam_attempts = debug_meta.get("sam_text_prompt_attempts", []) or []
    sam_prompt_audit = debug_meta.get("sam_prompt_selection_audit", []) or []
    sam_scores = debug_meta.get("sam_mask_scores", []) or []
    rag_blocks = debug_meta.get("rag_block_provenance", []) or []
    rag_summary = participation.get("rag_source_backed_summary", {}) or {}

    crop_items = []
    for item in crop_audit[:4]:
        if not isinstance(item, dict):
            continue
        label = item.get("label") or item.get("source") or "crop"
        source = item.get("source", "unknown")
        reason = item.get("reason") or item.get("status") or ""
        score = item.get("heatmap_score", item.get("sam_score", ""))
        score_text = f", score={float(score):.3f}" if isinstance(score, (int, float)) else ""
        crop_items.append(
            f"<b>{_html(label)}</b> from <code>{_html(source)}</code> "
            f"<span>({_html(reason)}{_html(score_text)})</span>"
        )

    prompt_items = []
    prompt_source = sam_prompt_audit if sam_prompt_audit else sam_attempts
    for item in prompt_source[:5]:
        if not isinstance(item, dict):
            continue
        role = item.get("role", "prompt")
        text = item.get("text", "")
        status = item.get("status", "attempted")
        reason = item.get("reason") or item.get("selection_reason") or item.get("source", "")
        prompt_items.append(
            f"<code>{_html(_display_status(role))}</code> {_html(text or 'box/ROI prompt')} "
            f"<span>-> {_html(_display_status(status))}{(' / ' + _html(_display_status(reason))) if reason else ''}</span>"
        )

    rag_regions = []
    graph_sources = []
    text_sources = []
    for block in rag_blocks:
        if not isinstance(block, dict):
            continue
        region = block.get("region")
        if region:
            rag_regions.append(str(region))
        graph_path = block.get("graph_cache_path")
        if graph_path:
            graph_sources.append(_source_label(graph_path))
        text_path = block.get("source_json_path")
        if text_path:
            text_sources.append(_source_label(text_path))

    phase_report = _as_text(debug_meta.get("phase_1_result")).strip() or "No Phase-1 report recorded."
    prompt_crop_labels = debug_meta.get("prompt_crop_labels", []) or []
    final_prompt_note = "built and sent" if final_prompt else "not generated yet"
    task_policy = debug_meta.get("task_policy", {}) or {}
    if isinstance(task_policy, dict):
        policy_bits = []
        for key, label in [
            ("use_local_anomaly_stream", "local anomaly stream"),
            ("use_mcts_search", "MCTS search"),
            ("use_sam3_local_refinement", "SAM3 refinement"),
            ("show_phase1_report_in_final_prompt", "Phase-1 report in Phase-2"),
            ("show_rag_blocks_in_final_prompt", "PVLA/RAG in Phase-2"),
        ]:
            if task_policy.get(key):
                policy_bits.append(label)
        phase1_role = _display_status(task_policy.get("phase1_logic_role") or "standard")
        policy_text = f"{', '.join(policy_bits) or 'global-only'}; phase1 role={phase1_role}"
    else:
        policy_text = _short_text(task_policy, 160) or "standard"

    mcts_actions = len(debug_meta.get("mcts_action_trace", []) or [])
    if isinstance(mcts_budget, dict) and mcts_budget:
        mcts_budget_text = (
            f"{mcts_budget.get('n_simulations', '?')} sims, depth "
            f"{mcts_budget.get('max_depth', '?')}, actions {mcts_budget.get('action_count', '?')}"
        )
    else:
        mcts_budget_text = _compact_json(mcts_budget, 120) or "not recorded"
    heatmap_peak = debug_meta.get("heatmap_peak_score", debug_meta.get("heatmap_score", 0))
    try:
        heatmap_peak_text = f"{float(heatmap_peak):.3f}"
    except (TypeError, ValueError):
        heatmap_peak_text = _as_text(heatmap_peak) or "n/a"
    selected_regions = []
    prompt_rag_audit = debug_meta.get("prompt_rag_selection_audit", {}) or {}
    if isinstance(prompt_rag_audit, dict):
        selected_regions.extend(prompt_rag_audit.get("selected_regions", []) or [])
    selected_regions.extend(rag_regions)
    sam_veto = debug_meta.get("sam3_normal_part_veto", {}) or {}
    if isinstance(sam_veto, dict) and sam_veto.get("status"):
        sam_gate = f"normal-part veto {_display_status(sam_veto.get('status'))}"
    elif int(debug_meta.get("sam_refined_crop_count", 0) or 0) > 0:
        sam_gate = f"{debug_meta.get('sam_refined_crop_count')} accepted refined crop(s)"
    elif sam_scores:
        sam_gate = "mask audited; no refined crop selected"
    else:
        sam_gate = "not needed for this task/category"

    global_panel = (
        '<section class="stream-panel global-panel">'
        '<div class="stream-heading"><div><p class="eyebrow">stream 1</p><h3>Global & logic stream</h3>'
        '<p>Whole-image reasoning supplies the semantic frame; SAM3 and logic gates check whether the local evidence is structurally valid.</p></div></div>'
        '<div class="chip-row">'
        + _chip_html("Phase-1", "report captured" if _as_text(debug_meta.get("phase_1_result")).strip() else "no report", "global")
        + _chip_html("policy", policy_text, "global")
        + _chip_html("SAM3", _display_status(participation.get("sam_participation_status", "unknown")), "global")
        + "</div>"
        f'<blockquote>{_html(_short_text(phase_report, 420))}</blockquote>'
        '<details><summary>Structural gate details</summary>'
        + _item_list_html([
            f"SAM3 gate: <code>{_html(sam_gate)}</code>",
            f"Mask scores: <code>{_html(_compact_json(sam_scores[:6], 180))}</code>",
            f"Artifact rule: <code>{_html(_short_text(debug_meta.get('sam3_crop_artifact_rule', 'not recorded'), 160))}</code>",
            *prompt_items,
        ], "No SAM3 prompt or mask audit was needed.")
        + "</details></section>"
    )

    mcts_status = _display_status(participation.get("mcts_participation_status", "unknown"))
    action_viz = _mcts_action_viz_html(mcts_search, debug_meta.get("mcts_action_trace", []), mcts_status)
    local_panel = (
        '<section class="stream-panel action-panel">'
        '<div class="stream-heading"><div><p class="eyebrow">stream 2</p><h3>Fine-grained & actions stream</h3>'
        '<p>AdaptCLIP proposes suspicious regions, MCTS searches compact crops, and only selected local views are sent forward.</p></div></div>'
        '<div class="chip-row">'
        + _chip_html("heatmap peak", heatmap_peak_text, "action")
        + _chip_html("MCTS", mcts_status, "action")
        + _chip_html("proposals", debug_meta.get("region_proposal_count", 0), "action")
        + _chip_html("crops", debug_meta.get("crop_count", 0), "action")
        + "</div>"
        + action_viz
        + '<details open><summary>Search summary</summary>'
        + _item_list_html([
            f"Localizer: <code>{_html(debug_meta.get('threshold_source', 'unknown'))}</code>",
            f"Proposals: <code>{_html(debug_meta.get('region_proposal_count', 0))}</code>, selected crops <code>{_html(debug_meta.get('crop_count', 0))}</code>, prompt-visible <code>{_html(debug_meta.get('prompt_visible_crop_count', 0))}</code>",
            f"MCTS budget: <code>{_html(mcts_budget_text)}</code>",
            *crop_items,
        ], "No local crop was selected for this question.")
        + "</details></section>"
    )

    fusion_panel = (
        '<section class="stream-panel fusion-panel">'
        '<div class="stream-heading"><div><p class="eyebrow">fusion</p><h3>Cross-stream evidence handoff</h3>'
        '<p>The final prompt is assembled from global facts, action crops, SAM3 gates, and graph-backed PVLA recall.</p></div></div>'
        '<div class="chip-row">'
        + _chip_html("prompt crops", len(prompt_crop_labels), "")
        + _chip_html("PVLA blocks", len(debug_meta.get("prompt_rag_text_blocks", []) or []), "pvla")
        + _chip_html("final prompt", final_prompt_note, "")
        + "</div>"
        + _item_list_html([
            f"Prompt crops: <code>{_html(', '.join(prompt_crop_labels) or 'none')}</code>",
            f"Selected regions: <code>{_html(', '.join(_first_items(selected_regions, 6)) or 'none')}</code>",
            f"Verification strategy: <code>{_html(debug_meta.get('verification_strategy', 'standard'))}</code>",
        ], "Phase-2 prompt has not been generated.")
        + "</section>"
    )

    verdict_panel = (
        '<section class="stream-panel verdict-panel">'
        '<div class="stream-heading"><div><p class="eyebrow">verdict</p><h3>Final verifier receives the curated bundle</h3>'
        '<p>The answer card below contains the model choice, ground truth, and binary AD override when applicable.</p></div></div>'
        '<div class="chip-row">'
        + _chip_html("MCTS", f"{_display_status(participation.get('mcts_participation_status', 'unknown'))} / {mcts_actions} actions", "action")
        + _chip_html("PVLA", f"{rag_summary.get('block_count', len(rag_blocks))} blocks", "pvla")
        + _chip_html("SAM3", f"{len(sam_scores)} masks", "global")
        + "</div></section>"
    )

    return (
        '<div class="dual-board-html">'
        '<div class="board-head"><div><p class="eyebrow">trace-backed demo</p>'
        '<h2>Dual-stream verification for this QA</h2>'
        '<p>Every panel is rendered from this run trace, pairing visuals with the text evidence that made them useful.</p>'
        '</div><span class="process-badge">Trace-backed</span></div>'
        '<div class="board-grid">'
        + _atlas_panel_html(category, rag_blocks)
        + global_panel
        + local_panel
        + fusion_panel
        + verdict_panel
        + "</div></div>"
    )

def _method_process_placeholder():
    return render_dual_stream_placeholder()

def _method_process_html(debug_meta, final_prompt="", category=""):
    return render_dual_stream_trace(debug_meta, final_prompt, category)

def _final_result_markdown(question, options, pred_key, gt_key, is_correct, final_response, debug_meta, ad_override):
    options = options or {}
    pred_text = options.get(pred_key, "") if pred_key else ""
    gt_text = options.get(gt_key, "") if gt_key else ""
    result_label = "Unknown"
    if is_correct is True:
        result_label = "Correct"
    elif is_correct is False:
        result_label = "Wrong"

    option_lines = []
    for key in sorted(options.keys()):
        marker = []
        if key == pred_key:
            marker.append("pred")
        if key == gt_key:
            marker.append("gt")
        suffix = f" ({', '.join(marker)})" if marker else ""
        option_lines.append(f"- **{key}.** {options[key]}{suffix}")

    gate_line = "not applicable"
    if ad_override:
        gate_line = (
            f"{'applied' if ad_override.get('applied') else 'checked'} | "
            f"{ad_override.get('reason', 'no reason')} | "
            f"{ad_override.get('original_pred_key')} -> {ad_override.get('override_pred_key')}"
        )

    return "\n".join([
        "### Question Result",
        f"**Question:** {question}",
        "",
        "| Prediction | Ground Truth | Result | Binary AD Gate |",
        "| --- | --- | --- | --- |",
        f"| **{pred_key or 'N/A'}** {pred_text} | **{gt_key or 'N/A'}** {gt_text} | **{result_label}** | {gate_line} |",
        "",
        "#### Options",
        "\n".join(option_lines) if option_lines else "No options loaded.",
        "",
        _method_trace_markdown(debug_meta),
        "",
        "#### Model Response",
        _as_text(final_response).strip() or "No response.",
    ])

async def run_analysis_stream(image, question, subclass, task_type, options, ground_truth=None, override_rag_text=None, override_rag_files=None):
    if not global_sys.is_initialized:
        yield None, "", None, None, None, "System not initialized.", "", "", None, None, None, "", _method_process_placeholder()
        return
    if image is None or not question:
        yield None, "", None, None, None, "No QA sample selected.", "", "", None, None, None, "", _method_process_placeholder()
        return
    
    # Debug Options Passing
    if not options:
        print("[DEBUG] No options received in run_analysis_stream.")
    else:
        print(f"[DEBUG] Options received: {len(options)} keys")

    # --- RAG Loading Logic (Updated for Smart Splitting) ---
    rag_blocks = []
    
    if override_rag_text or (override_rag_files is not None and len(override_rag_files) > 0):
        print("Using manual RAG context overrides (auto-splitting global/local).")
        
        # 1. Bucket sorting for images based on filename
        global_imgs = []
        local_imgs = []
        
        if override_rag_files:
            for item in override_rag_files:
                path = None
                # Extract path safely
                if isinstance(item, (list, tuple)) and len(item) > 0:
                    path = item[0]
                elif isinstance(item, dict):
                    path = item.get('name') or item.get('path')
                elif hasattr(item, 'name'):
                    path = item.name
                elif isinstance(item, str):
                    path = item
                
                if path and os.path.exists(path):
                    filename = os.path.basename(path).lower()
                    # Backend Logic: Checks if 'whole' or 'global' is in region name
                    if "whole" in filename or "global" in filename:
                        global_imgs.append(path)
                    else:
                        local_imgs.append(path)

        # 2. Construct Blocks
        # Block A: Global/Whole (Targeted by Phase 1)
        # We attach text here so Phase 1 knows what "Normal" looks like
        rag_blocks.append({
            "region": "Manual Global Reference (Whole Object)", 
            "text": override_rag_text if override_rag_text else "",
            "images": global_imgs
        })

        # Block B: Local/Regions (Targeted by Phase 2)
        # We DUPLICATE text here so Phase 2 knows "Defect" definitions
        # (Phase 2 explicitly skips blocks named 'whole', so it needs its own copy)
        rag_blocks.append({
            "region": "Manual Local Defect Reference",
            "text": override_rag_text if override_rag_text else "",
            "images": local_imgs
        })
    else:
        rag_blocks = global_sys.rag_agent.get_rag_context(subclass)

    # Prepare Display Data
    rag_file_paths, rag_text_content = [], ""
    if rag_blocks:
        for i, b in enumerate(rag_blocks):
            t_content = b.get('text', '').strip()
            if t_content:
                rag_text_content += f"=== {b.get('region', f'Region {i+1}')} ===\n{t_content}\n\n"
            
            if b.get('images'):
                for img_path in b['images']:
                    if os.path.exists(img_path):
                        rag_file_paths.append(img_path)
    
    state = {"heatmap": None, "log": "Starting GLLS method trace."}
    if override_rag_text:
        state["log"] += "\nNOTE: Running with edited knowledge blocks."
    if global_sys.vlm.use_vllm:
        state["log"] += "\nvLLM acceleration: active."

    configure_support = getattr(global_sys.localizer, "configure_support", None)
    if callable(configure_support):
        dataset_root = global_sys.data_manager._resolve_dataset_root()
        support_paths = find_normal_support_images(
            dataset_root,
            global_sys.data_manager.dataset_name,
            subclass,
            getattr(global_sys.args, "k_shot", 1),
        )
        configure_support(subclass, support_paths)
        state["log"] += f"\nAdaptCLIP support: {len(support_paths)} normal reference image(s)."

    # Yield 1: Init
    yield rag_file_paths, rag_text_content, None, None, None, state["log"], "", "", None, None, None, "", _method_process_placeholder()

    async def callback(key, value):
        if key == "heatmap":
            state["heatmap"] = value
        if key == "log":
            state["log"] += f"\n{value}"

    row = {'image': image.convert("RGB"), 'question': question, 'options': options, 'category': subclass, 'type': task_type}
    try:
        mcts = StreamableMCTS(row, global_sys.args, global_sys.vlm, global_sys.localizer, global_sys.rag_agent, rag_blocks, global_sys.sam_engine, callback)
        mcts_task = asyncio.create_task(mcts.process())
        
        while not mcts_task.done():
            await asyncio.sleep(0.1)
            yield rag_file_paths, rag_text_content, state["heatmap"], None, None, state["log"], "", "", None, None, None, "", _method_process_placeholder()
        
        result = await mcts_task
        
        # Save crops for gallery
        temp_dir = tempfile.mkdtemp()
        crop_file_paths = []
        gallery_items = []
        
        if result.get("red_box_image"):
            rb_path = os.path.join(temp_dir, "global_trace.png")
            result["red_box_image"].save(rb_path)
            crop_file_paths.append(rb_path)
            gallery_items.append((result["red_box_image"], "Global Scope"))
            
        for i, crop in enumerate(result.get("crop_images", [])):
            c_path = os.path.join(temp_dir, f"focus_crop_{i}.png")
            crop.save(c_path)
            crop_file_paths.append(c_path)
            gallery_items.append((crop, f"Focus View {i+1}"))
        
        debug_info = result.get("logic_debug_info", {})
        debug_view_img = debug_info.get("view_image")
        raw_p1_prompt = debug_info.get("phase1_prompt", "")
        
        formatted_prompt = raw_p1_prompt
        if formatted_prompt:
            formatted_prompt = formatted_prompt.replace("=== PHASE 1: GLOBAL INSPECTION ===", "### PHASE 1: GLOBAL INSPECTION")
            formatted_prompt = formatted_prompt.replace("STEP 1:", "\n**STEP 1:**").replace("STEP 2:", "\n**STEP 2:**").replace("STEP 3:", "\n**STEP 3:**")
            formatted_prompt = formatted_prompt.replace("INSTRUCTION:", "\n#### INSTRUCTION:")
            formatted_prompt = formatted_prompt.replace("[SPECIAL INSPECTION VIEW]", "**[SPECIAL INSPECTION VIEW]**")
        else:
            formatted_prompt = "*No prompt data available yet.*"
        
        # Build Final Prompt
        final_content = []
        if result.get("rag_content_list"):
            final_content.extend(result["rag_content_list"])
        if result.get("red_box_image"):
            final_content.append({"type": "text", "text": "Image: Global Trace (Red Box)\n"})
            final_content.append({"type": "image", "image": result["red_box_image"]})
        for i, crop in enumerate(result.get("crop_images", [])):
            final_content.append({"type": "text", "text": f"Image: Local Refinement {i+1}\n"})
            final_content.append({"type": "image", "image": crop})
        
        final_prompt = result.get("prompt")
        final_content.append({"type": "text", "text": final_prompt})
        
        debug_meta = result.get("debug_metadata", {}) or {}
        state["log"] += "\nGenerating final diagnosis."
        final_response = global_sys.vlm.generate(final_content)
        pred_key = _parse_answer_key(final_response)
        gt_key = _as_text(ground_truth).strip()
        ad_override = binary_anomaly_decision_override(task_type, options, debug_meta, pred_key) if pred_key else None
        if ad_override:
            pred_key = ad_override.get("override_pred_key", pred_key)
            debug_meta["binary_anomaly_decision_override"] = ad_override
        is_correct = (pred_key == gt_key) if pred_key and gt_key else None
        final_result_md = _final_result_markdown(
            question,
            options,
            pred_key,
            gt_key,
            is_correct,
            final_response,
            debug_meta,
            ad_override,
        )
        method_process_md = _method_process_html(debug_meta, final_prompt, subclass)
        state["log"] += (
            "\nAnalysis complete."
            f"\nPrediction: {pred_key or 'N/A'} | GT: {gt_key or 'N/A'} | Result: "
            f"{'Correct' if is_correct is True else 'Wrong' if is_correct is False else 'Unknown'}"
        )
        
        yield (
            rag_file_paths,     
            rag_text_content,   
            state["heatmap"],   
            gallery_items,      
            crop_file_paths,    
            state["log"],       
            final_result_md,     
            final_prompt,       
            result.get("red_box_image"), 
            crop_file_paths,
            debug_view_img,
            formatted_prompt,
            method_process_md
        )
    except Exception as e:
        import traceback
        traceback.print_exc()
        yield rag_file_paths, rag_text_content, state["heatmap"], None, None, f"Error: {str(e)}", "", "", None, None, None, "", f"<div class='method-process-card'><h3>Method run failed</h3><p>{_html(str(e))}</p></div>"

# Wrapper for Gradio Generator
def _blocked_run_outputs(message):
    return (
        None,
        "",
        None,
        None,
        None,
        message,
        f"### Runtime configuration required\n\n{message}",
        "",
        None,
        None,
        None,
        "",
        _runtime_config_error_html(message),
    )


def _weld_rule_markdown(result):
    global_result = result.get("global") or {}
    rows = [
        f"| `{rule.get('name', 'rule')}` | {rule.get('criterion', '')} | **{rule.get('status', 'unknown')}** |"
        for rule in global_result.get("rules") or []
    ]
    return "\n".join([
        "### Stream 1 · Global weld logic",
        "",
        f"Mask source: `{global_result.get('mask_source', 'unknown')}`  ",
        f"Station: `{global_result.get('station', 'unknown')}`  ",
        f"SAM3 quality: `{float(global_result.get('sam3_quality', 0.0) or 0.0):.3f}`",
        "",
        "| Rule | Criterion | Status |",
        "| --- | --- | --- |",
        *(rows or ["| No rule result | — | — |"]),
        "",
        "Insufficient structural evidence remains non-triggering; it is not promoted to a confident NG.",
    ])


def _weld_final_markdown(result, question):
    trace = result.get("trace") or {}
    local = trace.get("local_stream") or {}
    verdict = str(result.get("verdict") or "unknown")
    reference = str(trace.get("reference_ground_truth") or "UNKNOWN").upper()
    if reference in {"OK", "NG"}:
        comparison = "MATCH" if verdict == reference else "MISS" if reference == "NG" else "FALSE ALARM"
    else:
        comparison = "UNSCORED"
    reason = (
        "global weld logic triggered"
        if result.get("global_logic_ng")
        else "local AdaptCLIP/PVLA score triggered"
        if result.get("local_ng")
        else "neither stream triggered"
    )
    return "\n".join([
        "### Gear-weld inspection result",
        "",
        f"**Question:** {question}",
        "",
        f"> **{verdict}** — {reason}.",
        "",
        f"**Reference:** `{reference}` · **Comparison:** `{comparison}`",
        "",
        "| Global logic | Compact hotspot | PVLA normal similarity | Local score / threshold |",
        "| --- | ---: | ---: | ---: |",
        (
            f"| {'NG' if result.get('global_logic_ng') else 'pass'} | "
            f"{float(local.get('compact_hotspot_score', 0.0) or 0.0):.4f} | "
            f"{float(local.get('pvla_normal_similarity', 0.0) or 0.0):.4f} | "
            f"{float(local.get('orchestra_score', 0.0) or 0.0):.4f} / "
            f"{float(local.get('alert_threshold', 0.0) or 0.0):.4f} |"
        ),
        "",
        "**Reference contract:** `dx` is an anomalous defect region; `ljtq` is an allowed station-2 "
        "connection context. The archive has no official benchmark split, so this UI does not claim aggregate accuracy.",
    ])


def _weld_result_outputs(result, question):
    trace = result.get("trace") or {}
    local = trace.get("local_stream") or {}
    references = global_sys.weld_runtime.normal_reference_paths()
    gallery = result.get("hotspot_gallery") or []
    crop_paths = [str(item[0]) for item in gallery if isinstance(item, (list, tuple)) and item]
    reference_text = "\n".join([
        "GLLS gear-weld knowledge route",
        "- Source rules: 焊缝缺陷检测升级.pdf",
        "- Offline atlas: 14 graph nodes / 18 graph edges",
        "- Visual prior: one trusted normal weld image and deterministic derivatives",
        "- Global evidence: SAM3-assisted geometry and weld continuity rules",
        "- Local evidence: overlapping high-resolution AdaptCLIP tiles",
        "- Fusion: global_logic_ng OR local_adaptclip_ng",
        "- Evaluation status: qualitative only; image-level GT unavailable",
        f"- PVLA reference patches: {(result.get('pvla') or {}).get('reference_count', 0)}",
    ])
    fusion_record = "\n".join([
        "GLLS weld fusion record",
        f"sample={trace.get('sample_id', 'unknown')}",
        f"global_logic_ng={bool(result.get('global_logic_ng'))}",
        f"local_adaptclip_ng={bool(result.get('local_ng'))}",
        f"local_score={float(local.get('orchestra_score', 0.0) or 0.0):.6f}",
        f"threshold={float(local.get('alert_threshold', 0.0) or 0.0):.6f}",
        f"verdict={result.get('verdict', 'unknown')}",
    ])
    return (
        references,
        reference_text,
        result.get("heatmap_view"),
        gallery,
        crop_paths,
        json.dumps(trace, ensure_ascii=False, indent=2),
        _weld_final_markdown(result, question),
        fusion_record,
        result.get("annotation_view"),
        crop_paths,
        (result.get("global") or {}).get("overlay"),
        _weld_rule_markdown(result),
        render_weld_trace(result),
    )


def _weld_running_outputs():
    try:
        references = global_sys.weld_runtime.normal_reference_paths()
    except Exception:
        references = []
    return (
        references,
        "Loading source-backed weld knowledge and running both evidence streams…",
        None,
        None,
        None,
        "Starting GLLS gear-weld inspection.",
        "### Gear-weld inspection in progress",
        "",
        None,
        None,
        None,
        "Running global weld logic…",
        render_weld_running(),
    )


def search_runner_wrapper(
    img,
    q,
    sub,
    task,
    opts,
    gt=None,
    dataset_name=None,
    localizer_choice="Auto",
    k_shot=1,
    ckpt_path=None,
    gpu_id=None,
    man_txt=None,
    man_files=None,
):
    runtime_ok, _, runtime_message = global_sys.runtime_match_status(
        dataset_name,
        localizer_choice,
        k_shot,
        ckpt_path,
        gpu_id,
    )
    if not runtime_ok:
        yield _blocked_run_outputs(runtime_message)
        return

    if str(dataset_name).lower() == WELD_DATASET:
        yield _weld_running_outputs()
        try:
            result = global_sys.weld_runtime.inspect(task)
            yield _weld_result_outputs(result, q)
        except Exception as e:
            import traceback
            traceback.print_exc()
            yield _blocked_run_outputs(f"Weld inspection failed: {e}")
        return

    # Verify opts before running
    opts = _coerce_options(opts)
    if not opts:
        print("[WRAPPER] Options are None or invalid, resetting to empty dict.")
        
    loop = asyncio.new_event_loop()
    asyncio.set_event_loop(loop)
    gen = run_analysis_stream(img, q, sub, task, opts, gt, man_txt, man_files)
    try:
        while True:
            yield loop.run_until_complete(gen.__anext__())
    except StopAsyncIteration:
        pass
    finally:
        loop.close()

# ==========================================
# 6. Gradio UI
# ==========================================

CSS = """
:root {
    --canvas: #F1F5FA;
    --panel: #FFFFFF;
    --panel-raised: #FBFCFE;
    --rule: #D8E1EC;
    --ink: #142033;
    --muted: #627087;
    --signal: #315BE8;
    --signal-hover: #2448C4;
    --signal-subtle: #EDF2FF;
    --builder: #B0640F;
    --builder-subtle: #F8EFE4;
    --pass: #0F7A4D;
    --pass-subtle: #EAF6F0;
    --danger: #C0392B;
    --radius-sm: 8px;
    --radius-md: 12px;
    --radius-lg: 20px;
    --shadow-sm: 0 1px 2px rgba(20, 32, 51, 0.04), 0 8px 24px rgba(20, 32, 51, 0.045);
    --shadow-md: 0 16px 38px rgba(20, 32, 51, 0.10), 0 2px 8px rgba(20, 32, 51, 0.05);
}
html, body, .gradio-container {
    background:
        radial-gradient(circle at 8% 5%, rgba(49, 91, 232, 0.075), transparent 28rem),
        radial-gradient(circle at 92% 13%, rgba(15, 122, 77, 0.055), transparent 30rem),
        linear-gradient(90deg, rgba(216, 225, 236, 0.46) 1px, transparent 1px) 0 0 / 88px 88px,
        var(--canvas) !important;
    color: var(--ink);
    font-family: Inter, "IBM Plex Sans", "Segoe UI", "Microsoft YaHei UI", system-ui, sans-serif;
}
.gradio-container {
    max-width: none !important;
}
.gradio-container .contain {
    width: 100% !important;
    max-width: none !important;
    padding-bottom: 36px;
}
.app-topbar {
    position: sticky;
    top: 0;
    z-index: 20;
    margin: -16px -16px 0;
    padding: 15px 28px;
    border-bottom: 1px solid rgba(135, 161, 216, 0.25);
    background:
        linear-gradient(120deg, rgba(49, 91, 232, 0.20), transparent 42%),
        linear-gradient(135deg, #111E33, #0C1728 72%, #10243C);
    box-shadow: 0 12px 30px rgba(10, 23, 42, 0.14);
    backdrop-filter: blur(16px);
}
.brand-block {
    display: flex;
    align-items: center;
    gap: 12px;
}
.brand-mark {
    position: relative;
    width: 40px;
    min-width: 40px;
    min-height: 40px;
    border: 1px solid rgba(140, 169, 255, 0.50);
    border-radius: 13px;
    background: rgba(255, 255, 255, 0.06);
    box-shadow: inset 0 0 0 7px rgba(49, 91, 232, 0.11), 0 0 22px rgba(49, 91, 232, 0.16);
}
.brand-mark::before,
.brand-mark::after {
    content: "";
    position: absolute;
    inset: 10px;
    border: 3px solid #45D2D9;
    border-radius: 50%;
}
.brand-mark::after {
    inset: 17px;
    border: 0;
    background: #FFFFFF;
    box-shadow: 0 0 9px rgba(69, 210, 217, 0.85);
}
.brand-copy strong {
    display: block;
    font-family: Literata, "Source Serif 4", Georgia, "Microsoft YaHei", serif;
    color: #FFFFFF;
    font-size: 21px;
    line-height: 1.05;
}
.brand-copy small {
    color: #B7C5DA;
}
.hero-copy,
.section-copy,
.muted-note {
    color: var(--muted);
}
.status-chip textarea,
.status-chip input {
    min-height: 40px !important;
    border-radius: 12px !important;
    border: 1px solid rgba(155, 177, 213, 0.32) !important;
    background: rgba(255, 255, 255, 0.08) !important;
    color: #E6EDF8 !important;
    font-family: "Cascadia Code", Consolas, monospace !important;
    font-size: 0.78rem !important;
}
.hero-panel {
    position: relative;
    display: grid;
    grid-template-columns: minmax(0, 1fr) minmax(260px, 360px);
    gap: 18px;
    align-items: end;
    overflow: hidden;
    margin: 20px 0 14px;
    padding: 24px 26px;
    border: 1px solid var(--rule);
    border-radius: var(--radius-lg);
    background:
        linear-gradient(105deg, rgba(49, 91, 232, 0.045), transparent 44%),
        var(--panel);
    box-shadow: var(--shadow-sm);
}
.hero-panel h1 {
    margin: 4px 0 8px;
    color: var(--ink);
    font-family: Literata, "Source Serif 4", Georgia, "Microsoft YaHei", serif;
    font-size: clamp(2rem, 3.4vw, 3.15rem);
    line-height: 1;
    letter-spacing: -0.025em;
}
.eyebrow {
    color: var(--signal);
    font-size: 0.72rem;
    font-weight: 900;
    letter-spacing: 0.10em;
    text-transform: uppercase;
}
.hero-metrics {
    display: grid;
    gap: 8px;
}
.hero-metrics span {
    display: flex;
    justify-content: space-between;
    gap: 18px;
    padding: 9px 12px;
    border: 1px solid var(--rule);
    border-radius: var(--radius-md);
    background: var(--panel-raised);
}
.hero-metrics b {
    color: var(--signal);
}
.custom-card {
    background: var(--panel);
    border-radius: var(--radius-lg);
    padding: 20px;
    border: 1px solid var(--rule);
    box-shadow: var(--shadow-sm);
}
.custom-card.tight {
    padding: 14px;
}
.runtime-control-panel {
    position: sticky;
    top: 72px;
    z-index: 18;
    margin: 12px 0 16px;
    padding: 14px 18px 12px;
    border: 1px solid var(--rule);
    border-radius: var(--radius-lg);
    background: rgba(255, 255, 255, 0.97);
    box-shadow: var(--shadow-md);
    backdrop-filter: blur(14px);
}
.runtime-action-bar {
    align-items: center;
    gap: 12px;
}
.runtime-heading {
    min-width: 300px;
}
.runtime-heading .section-title {
    margin-bottom: 2px;
}
.runtime-heading .section-copy {
    margin-bottom: 0;
}
.runtime-action button {
    min-height: 48px !important;
}
.runtime-select-grid {
    align-items: end;
    gap: 12px;
    margin-top: 6px;
}
.runtime-select-grid > div {
    min-width: 160px;
}
.sample-actions button {
    min-height: 44px !important;
    margin-top: 2px;
}
.demo-shell {
    align-items: flex-start;
    gap: 14px;
}
.dual-board-card {
    background: var(--panel);
    border: 1px solid var(--rule);
    border-radius: var(--radius-lg);
    padding: 20px;
    box-shadow: var(--shadow-sm);
}
.dual-board-html {
    color: var(--ink);
}
.board-head {
    display: flex;
    justify-content: space-between;
    gap: 18px;
    align-items: flex-start;
    padding: 4px 2px 16px;
}
.board-head h2 {
    margin: 2px 0 6px;
    color: var(--ink);
    font-family: Literata, "Source Serif 4", Georgia, "Microsoft YaHei", serif;
    font-size: 1.55rem;
    line-height: 1.08;
}
.board-head p {
    margin: 0;
    color: var(--muted);
}
.board-grid {
    display: grid;
    grid-template-columns: repeat(2, minmax(0, 1fr));
    gap: 14px;
}
.weld-running-board {
    min-height: 390px;
    padding: 8px 4px;
}
.running-badge {
    border-color: #E6C38F;
    background: #FFF7E8;
    color: #9A5A12;
}
.running-badge i {
    width: 12px;
    height: 12px;
    margin-right: 7px;
    border: 2px solid rgba(154, 90, 18, 0.28);
    border-top-color: #9A5A12;
    border-radius: 50%;
    animation: weld-spin 0.85s linear infinite;
}
.run-progress-grid {
    display: grid;
    grid-template-columns: repeat(2, minmax(0, 1fr));
    gap: 12px;
    margin-top: 10px;
}
.run-step {
    display: flex;
    align-items: flex-start;
    gap: 12px;
    min-height: 96px;
    padding: 14px;
    border: 1px solid var(--rule);
    border-radius: var(--radius-md);
    background: var(--panel-raised);
}
.run-step > b {
    display: grid;
    place-items: center;
    flex: 0 0 30px;
    width: 30px;
    height: 30px;
    border-radius: 50%;
    background: var(--signal-subtle);
    color: var(--signal);
}
.run-step span,
.run-step strong,
.run-step em {
    display: block;
}
.run-step strong {
    color: var(--ink);
    font-size: 0.94rem;
}
.run-step em {
    margin-top: 4px;
    color: var(--muted);
    font-size: 0.8rem;
    font-style: normal;
    line-height: 1.4;
}
.run-wait-note {
    margin-top: 14px !important;
    padding: 10px 12px;
    border-left: 4px solid var(--builder);
    border-radius: 0 var(--radius-sm) var(--radius-sm) 0;
    background: var(--builder-subtle);
    color: var(--ink) !important;
    font-weight: 700;
}
@keyframes weld-spin {
    to { transform: rotate(360deg); }
}
.stream-panel {
    border: 1px solid var(--rule);
    border-radius: var(--radius-md);
    background: var(--panel-raised);
    padding: 14px;
    min-height: 190px;
}
.stream-panel h3 {
    margin: 0 0 6px;
    color: var(--ink);
    font-size: 1.05rem;
}
.stream-panel p {
    margin: 0 0 8px;
    color: var(--muted);
}
.stream-panel blockquote {
    margin: 12px 0 0;
    padding: 10px 12px;
    border-left: 4px solid var(--signal);
    border-radius: 0 var(--radius-sm) var(--radius-sm) 0;
    background: var(--signal-subtle);
    color: var(--ink);
}
.stream-panel details {
    margin-top: 10px;
    border-top: 1px solid var(--rule);
    padding-top: 9px;
}
.stream-panel summary {
    cursor: pointer;
    color: var(--ink);
    font-weight: 800;
}
.atlas-panel,
.pvla-recall-panel,
.online-crop-panel,
.fusion-panel,
.verdict-panel {
    grid-column: 1 / -1;
}
.global-panel {
    border-top: 4px solid var(--signal);
}
.action-panel {
    border-top: 4px solid var(--builder);
}
.atlas-panel {
    border-top: 4px solid var(--pass);
}
.pvla-recall-panel {
    border-top: 4px solid #2F8A6A;
}
.online-crop-panel {
    border-top: 4px solid var(--builder);
}
.fusion-panel {
    border-top: 4px solid #5A6675;
}
.verdict-panel {
    border-top: 4px solid #8A5A22;
}
.chip-row {
    display: flex;
    flex-wrap: wrap;
    gap: 8px;
    margin: 10px 0;
}
.evidence-chip {
    display: inline-flex;
    align-items: baseline;
    gap: 6px;
    max-width: 100%;
    padding: 6px 9px;
    border: 1px solid var(--rule);
    border-radius: 999px;
    background: #F8FAFD;
    color: var(--ink);
    font-size: 0.8rem;
}
.evidence-chip b {
    color: var(--muted);
    font-weight: 800;
}
.evidence-chip em {
    overflow-wrap: anywhere;
    font-style: normal;
    font-weight: 800;
}
.evidence-chip.global {
    background: var(--signal-subtle);
    border-color: #CBD6F6;
    color: var(--signal);
}
.evidence-chip.action {
    background: var(--builder-subtle);
    border-color: #E9D2B6;
    color: var(--builder);
}
.evidence-chip.pvla {
    background: var(--pass-subtle);
    border-color: #C7E5D5;
    color: var(--pass);
}
.paired-grid {
    display: grid;
    grid-template-columns: minmax(0, 1fr) minmax(0, 1.2fr);
    gap: 12px;
}
.graph-mini {
    position: relative;
    min-height: 150px;
    padding: 12px;
    border: 1px solid #C7E5D5;
    border-radius: var(--radius-md);
    background:
        linear-gradient(90deg, rgba(15, 122, 77, 0.12) 1px, transparent 1px) 0 0 / 34px 34px,
        var(--pass-subtle);
}
.graph-mini strong {
    display: inline-block;
    margin-bottom: 8px;
    color: var(--pass);
    font-size: 1.02rem;
}
.graph-mini ul,
.stream-panel ul {
    margin: 8px 0 0;
    padding-left: 18px;
}
.graph-mini li,
.stream-panel li {
    margin: 5px 0;
    overflow-wrap: anywhere;
}
.weld-crop-grid {
    display: grid;
    grid-template-columns: repeat(4, minmax(0, 1fr));
    gap: 10px;
    margin-top: 10px;
}
.weld-evidence-figure {
    margin: 10px 0 0;
    padding: 8px;
    border: 1px solid var(--rule);
    border-radius: var(--radius-md);
    background: var(--panel-raised);
}
.weld-evidence-figure img {
    display: block;
    width: 100%;
    max-height: 420px;
    object-fit: contain;
    border-radius: calc(var(--radius-md) - 3px);
    background: #0F1720;
}
.weld-evidence-figure figcaption {
    margin-top: 7px;
    color: var(--muted);
    font-size: 0.78rem;
    font-weight: 750;
    overflow-wrap: anywhere;
}
.atlas-graph-wrap {
    position: relative;
    max-width: 980px;
    margin: 10px auto 0;
    padding: 14px 14px 4px;
    border: 1px solid #C7E5D5;
    border-radius: var(--radius-md);
    background:
        radial-gradient(circle at 50% 50%, rgba(15, 122, 77, 0.10), transparent 40%),
        linear-gradient(90deg, rgba(15, 122, 77, 0.06) 1px, transparent 1px) 0 0 / 40px 40px,
        linear-gradient(0deg, rgba(15, 122, 77, 0.05) 1px, transparent 1px) 0 0 / 40px 40px,
        var(--pass-subtle);
    overflow: hidden;
}
.atlas-graph-svg {
    width: 100%;
    height: auto;
    display: block;
    overflow: visible;
}
.atlas-ring {
    fill: none;
    stroke: rgba(15, 122, 77, 0.20);
    stroke-width: 1.4;
    stroke-dasharray: 3 8;
}
.atlas-link {
    fill: none;
    stroke: #9FB7AC;
    stroke-width: 1.8;
    stroke-linecap: round;
    opacity: 0.9;
}
.atlas-link.hier {
    marker-end: url(#atlas-arrow);
}
.atlas-link.diff {
    stroke: #C98A57;
    stroke-width: 1.5;
    stroke-dasharray: 5 5;
    opacity: 0.82;
}
.atlas-graph-svg marker path {
    fill: #8AA79A;
}
.atlas-legend {
    display: flex;
    flex-wrap: wrap;
    gap: 14px;
    margin: 0 2px 6px;
    font-size: 0.74rem;
    color: var(--muted);
}
.atlas-legend .atlas-leg {
    display: inline-flex;
    align-items: center;
    gap: 6px;
    font-weight: 700;
}
.atlas-legend .atlas-leg i {
    width: 11px;
    height: 11px;
    border-radius: 50%;
    background: #8895A6;
    box-shadow: inset 0 0 0 1px rgba(0, 0, 0, 0.08);
}
.atlas-legend .atlas-leg.root i { background: #D7466B; }
.atlas-legend .atlas-leg.region i { background: #E0931F; }
.atlas-legend .atlas-leg.defect i { background: #3452C7; }
.atlas-legend .atlas-leg.diff i {
    width: 16px;
    height: 0;
    border-radius: 0;
    border-top: 2px dashed #C98A57;
    background: none;
    box-shadow: none;
}
.atlas-overflow {
    margin: 4px 2px 2px;
    color: var(--muted);
    font-size: 0.75rem;
    text-align: center;
}
.atlas-node {
    height: 100%;
    box-sizing: border-box;
    display: flex;
    align-items: center;
    gap: 7px;
    padding: 4px 11px;
    border: 1px solid rgba(78, 98, 118, 0.28);
    border-radius: 999px;
    background: rgba(255, 255, 255, 0.97);
    color: var(--ink);
    box-shadow: 0 4px 12px rgba(31, 43, 58, 0.12);
    overflow: hidden;
}
.atlas-node .atlas-dot {
    flex: 0 0 auto;
    width: 10px;
    height: 10px;
    border-radius: 50%;
    background: #8895A6;
}
.atlas-node .atlas-node-text {
    display: flex;
    flex-direction: column;
    min-width: 0;
    line-height: 1.05;
}
.atlas-node .atlas-node-text em {
    font-style: normal;
    font-size: 0.55rem;
    font-weight: 800;
    letter-spacing: 0.04em;
    text-transform: uppercase;
    color: var(--muted);
}
.atlas-node .atlas-node-text strong {
    font-size: 0.68rem;
    font-weight: 700;
    color: var(--ink);
    overflow: hidden;
    display: -webkit-box;
    -webkit-line-clamp: 2;
    -webkit-box-orient: vertical;
    word-break: break-word;
}
.atlas-node.root {
    border-color: rgba(215, 70, 107, 0.55);
    background: linear-gradient(180deg, #FFF1F4, #FFE3EA);
    box-shadow: 0 6px 18px rgba(215, 70, 107, 0.22);
}
.atlas-node.root .atlas-dot { width: 12px; height: 12px; background: #D7466B; }
.atlas-node.root .atlas-node-text em { color: #C76286; }
.atlas-node.root .atlas-node-text strong {
    font-size: 0.84rem;
    color: #B12E52;
    -webkit-line-clamp: 1;
}
.atlas-node.region {
    border-color: rgba(224, 147, 31, 0.5);
    background: #FFF7EA;
}
.atlas-node.region .atlas-dot { background: #E0931F; }
.atlas-node.region .atlas-node-text strong { color: #9A6310; }
.atlas-node.source {
    border-color: rgba(15, 122, 77, 0.5);
    background: #EFFAF3;
}
.atlas-node.source .atlas-dot { background: #0F7A4D; }
.atlas-node.source .atlas-node-text strong { color: #0C6A43; }
.atlas-node.defect {
    border-color: rgba(52, 82, 199, 0.45);
    background: #F1F4FF;
}
.atlas-node.defect .atlas-dot { background: #3452C7; }
.atlas-node.defect .atlas-node-text strong { color: #2A43A8; }
.atlas-crop-grid {
    display: grid;
    grid-template-columns: repeat(auto-fill, minmax(80px, 1fr));
    gap: 8px;
}
.atlas-crop {
    margin: 0;
}
.atlas-crop-img {
    aspect-ratio: 1 / 1;
    display: flex;
    align-items: center;
    justify-content: center;
    border: 1px solid var(--rule);
    border-radius: var(--radius-sm);
    background: #fff;
    overflow: hidden;
}
.atlas-crop-img img {
    width: 100%;
    height: 100%;
    object-fit: cover;
    display: block;
}
.atlas-crop-missing {
    padding: 4px;
    color: var(--muted);
    font-size: 0.64rem;
    text-align: center;
}
.atlas-crop figcaption {
    margin-top: 4px;
    color: var(--ink);
    font-size: 0.68rem;
    font-weight: 700;
    line-height: 1.15;
    word-break: break-word;
}
.atlas-crop-note,
.atlas-crop-empty {
    margin: 8px 0 0;
    color: var(--muted);
    font-size: 0.72rem;
}
.mcts-viz {
    display: grid;
    gap: 7px;
    margin: 8px 0 2px;
    padding: 10px 12px;
    border: 1px solid #E9D2B6;
    border-radius: var(--radius-md);
    background: var(--builder-subtle);
}
.mcts-viz.empty {
    color: var(--muted);
    font-size: 0.82rem;
    background: var(--panel-raised);
    border-color: var(--rule);
}
.mcts-act {
    display: grid;
    grid-template-columns: 104px 1fr;
    align-items: center;
    gap: 10px;
}
.mcts-act-h {
    display: inline-flex;
    align-items: center;
    gap: 6px;
    font-size: 0.78rem;
    font-weight: 800;
    color: var(--builder);
}
.mcts-act-h i {
    font-style: normal;
    font-size: 0.9rem;
}
.mcts-act-h b {
    margin-left: auto;
    color: var(--ink);
}
.mcts-bar {
    position: relative;
    height: 8px;
    border-radius: 999px;
    background: rgba(176, 100, 15, 0.16);
    overflow: hidden;
}
.mcts-bar span {
    position: absolute;
    inset: 0 auto 0 0;
    height: 100%;
    min-width: 2px;
    border-radius: 999px;
    background: var(--builder);
}
.mcts-act.zoom .mcts-bar span { background: var(--signal); }
.mcts-act.stop .mcts-bar span { background: var(--muted); }
.mcts-viz-meta {
    margin-top: 2px;
    color: var(--muted);
    font-size: 0.72rem;
}
.atlas-details {
    margin-top: 10px;
    border: 1px solid var(--rule);
    border-radius: var(--radius-md);
    background: rgba(255, 255, 255, 0.74);
}
.atlas-details summary {
    padding: 10px 12px;
    cursor: pointer;
    color: var(--ink);
    font-weight: 800;
}
.atlas-detail-grid {
    display: grid;
    grid-template-columns: minmax(0, 1fr) minmax(0, 1fr) minmax(0, 1fr);
    gap: 12px;
    padding: 0 12px 12px;
}
.atlas-detail-grid h4 {
    margin: 0 0 8px;
    color: var(--muted);
    font-size: 0.82rem;
    text-transform: uppercase;
}
.atlas-detail-grid ul {
    margin: 0;
    padding-left: 18px;
}
.atlas-detail-grid li {
    margin: 6px 0;
}
.atlas-detail-grid li span,
.atlas-detail-grid li em {
    color: var(--muted);
    font-style: normal;
}
.atlas-detail-grid p {
    margin: 0 0 8px;
    padding: 8px;
    border: 1px solid var(--rule);
    border-radius: var(--radius-sm);
    background: var(--panel);
}
.evidence-note {
    display: grid;
    grid-template-columns: repeat(2, minmax(0, 1fr));
    gap: 10px;
}
.evidence-note p {
    min-height: 74px;
    margin: 0;
    padding: 10px;
    border: 1px solid var(--rule);
    border-radius: var(--radius-sm);
    background: var(--panel);
    color: var(--ink);
}
.evidence-note b {
    color: var(--muted);
}
.evidence-pair {
    align-items: stretch;
}
.evidence-detail {
    min-height: 100%;
}
.developer-trace textarea {
    min-height: 220px !important;
}
.section-title {
    color: var(--ink);
    font-weight: 800;
    font-size: 1.08rem;
    border-left: 4px solid var(--signal);
    padding-left: 12px;
    margin-bottom: 8px;
}
.section-title-sm {
    color: var(--ink);
    font-weight: 800;
    font-size: 0.98rem;
    margin-bottom: 8px;
}
.section-copy {
    margin: 0 0 12px;
    font-size: 0.92rem;
}
.desk-btn-primary {
    background: linear-gradient(135deg, var(--signal), #416EF3) !important;
    color: white !important;
    font-weight: 800 !important;
    border-radius: var(--radius-md) !important;
    border: 1px solid var(--signal-hover) !important;
    box-shadow: 0 8px 20px rgba(49, 91, 232, 0.22), inset 0 1px 0 rgba(255, 255, 255, 0.18) !important;
    transition: transform 0.18s ease, box-shadow 0.18s ease !important;
}
.desk-btn-primary:hover {
    background: var(--signal-hover) !important;
    transform: translateY(-1px);
    box-shadow: 0 11px 24px rgba(49, 91, 232, 0.28) !important;
}
.desk-btn-secondary {
    background: var(--panel) !important;
    color: var(--ink) !important;
    font-weight: 700 !important;
    border: 1px solid var(--rule) !important;
    border-radius: var(--radius-md) !important;
}
.gradio-container button {
    border-radius: var(--radius-md) !important;
}
.gradio-container button:focus-visible,
.gradio-container input:focus-visible,
.gradio-container textarea:focus-visible,
.gradio-container select:focus-visible,
.gradio-container summary:focus-visible {
    outline: 3px solid rgba(49, 91, 232, 0.24) !important;
    outline-offset: 2px;
}
.gradio-container input,
.gradio-container textarea,
.gradio-container select {
    border-radius: var(--radius-sm) !important;
}
.gradio-container label,
.gradio-container .label-wrap span {
    color: var(--muted) !important;
    font-weight: 700 !important;
}
.desk-image-upload,
.evidence-image {
    border-radius: var(--radius-md);
    overflow: hidden;
}
.log-box textarea {
    background: #1B2230 !important;
    color: #B9F6D6 !important;
    font-family: "Cascadia Code", Consolas, monospace !important;
    border-radius: var(--radius-md) !important;
}
.prompt-card {
    background: var(--panel-raised);
    border: 1px solid var(--rule);
    border-radius: var(--radius-md);
    padding: 16px;
    min-height: 360px !important;
    max-height: 460px !important;
    overflow-y: auto;
    font-size: 0.95rem;
    line-height: 1.6;
}
.prompt-card h3 {
    color: var(--signal);
    border-bottom: 1px solid var(--rule);
    padding-bottom: 8px;
    margin-top: 10px;
}
.prompt-card strong {
    color: var(--ink);
    background-color: var(--signal-subtle);
    padding: 0 4px;
    border-radius: 4px;
}
.prompt-card h3 {
    font-family: Literata, "Source Serif 4", Georgia, "Microsoft YaHei", serif;
    font-size: 1.45rem;
}
.prompt-card h4 {
    margin: 16px 0 8px;
    color: var(--ink);
    border-top: 1px solid var(--rule);
    padding-top: 12px;
}
.prompt-card code {
    background: #EEF1FB;
    color: #1B2230;
    border: 1px solid #D8E0F2;
    border-radius: 4px;
    padding: 1px 4px;
}
.method-process-card {
    background: transparent;
    border: 0;
    padding: 0;
    font-size: 0.95rem;
    line-height: 1.55;
}
.method-process-card h3 {
    font-family: Literata, "Source Serif 4", Georgia, "Microsoft YaHei", serif;
    font-size: 1.45rem;
    margin: 0;
}
.method-process-card p {
    margin: 6px 0 0;
    color: var(--muted);
}
.process-head {
    display: flex;
    justify-content: space-between;
    align-items: flex-start;
    gap: 18px;
    padding: 2px 2px 14px;
}
.process-badge {
    display: inline-flex;
    align-items: center;
    min-height: 30px;
    padding: 4px 10px;
    border: 1px solid #B7D7C8;
    border-radius: 999px;
    background: var(--pass-subtle);
    color: var(--pass);
    font-size: 0.82rem;
    font-weight: 800;
    white-space: nowrap;
}
.method-kpi-grid {
    display: grid;
    grid-template-columns: repeat(4, minmax(0, 1fr));
    gap: 10px;
    margin: 0 0 12px;
}
.method-kpi {
    min-height: 68px;
    border: 1px solid var(--rule);
    border-radius: var(--radius-md);
    background: var(--panel-raised);
    padding: 10px 12px;
}
.method-kpi span {
    display: block;
    color: var(--muted);
    font-size: 0.78rem;
    font-weight: 800;
}
.method-kpi strong {
    display: block;
    margin-top: 4px;
    color: var(--ink);
    font-size: 1rem;
    overflow-wrap: anywhere;
}
.method-kpi.signal strong { color: var(--signal); }
.method-kpi.builder strong { color: var(--builder); }
.method-kpi.pass strong { color: var(--pass); }
.method-stage-grid {
    display: grid;
    grid-template-columns: repeat(2, minmax(0, 1fr));
    gap: 12px;
}
.method-stage {
    border: 1px solid var(--rule);
    border-radius: var(--radius-md);
    background: var(--panel-raised);
    padding: 14px;
    min-height: 220px;
}
.method-stage:nth-child(5),
.method-stage.phase-final {
    grid-column: 1 / -1;
    min-height: auto;
}
.method-stage.muted {
    min-height: 150px;
}
.stage-head {
    display: flex;
    gap: 10px;
    align-items: center;
}
.stage-head span {
    display: inline-grid;
    place-items: center;
    width: 32px;
    height: 32px;
    border-radius: 999px;
    background: var(--signal-subtle);
    color: var(--signal);
    font-weight: 900;
    font-size: 0.82rem;
}
.phase-local .stage-head span { background: var(--builder-subtle); color: var(--builder); }
.phase-sam .stage-head span { background: var(--pass-subtle); color: var(--pass); }
.phase-pvla .stage-head span { background: #F2F4F7; color: var(--ink); }
.phase-final .stage-head span { background: #F7F1EA; color: var(--builder); }
.method-stage h4 {
    margin: 0;
    color: var(--ink);
    font-size: 1rem;
    line-height: 1.25;
}
.method-stage ul {
    margin: 10px 0 0;
    padding-left: 18px;
}
.method-stage li {
    margin: 6px 0;
    color: var(--ink);
    overflow-wrap: anywhere;
}
.method-stage li span {
    color: var(--muted);
}
.method-process-card code {
    background: #EEF1FB;
    color: #1B2230;
    border: 1px solid #D8E0F2;
    border-radius: 4px;
    padding: 1px 4px;
    overflow-wrap: anywhere;
}
.logic-view-img {
    height: 360px !important;
    border-radius: var(--radius-md);
    overflow: hidden;
    border: 1px solid var(--rule);
    display: flex;
    justify-content: center;
    align-items: center;
    background: var(--panel-raised);
}
.hidden-artifact {
    display: none !important;
}
@media (max-width: 1100px) {
    .hero-panel {
        grid-template-columns: 1fr;
    }
    .paired-grid,
    .atlas-detail-grid,
    .evidence-note,
    .board-grid {
        grid-template-columns: 1fr;
    }
    .weld-crop-grid {
        grid-template-columns: repeat(2, minmax(0, 1fr));
    }
    .run-progress-grid {
        grid-template-columns: 1fr;
    }
    .method-kpi-grid,
    .method-stage-grid {
        grid-template-columns: 1fr;
    }
    .atlas-panel,
    .pvla-recall-panel,
    .online-crop-panel,
    .fusion-panel,
    .verdict-panel {
        grid-column: auto;
    }
    .method-stage:nth-child(5),
    .method-stage.phase-final {
        grid-column: auto;
    }
    .runtime-control-panel {
        position: relative;
        top: auto;
    }
    .runtime-heading {
        min-width: 0;
    }
    .runtime-action-bar,
    .runtime-select-grid {
        flex-wrap: wrap;
    }
}
@media (max-width: 900px) {
    .demo-shell,
    .evidence-pair {
        flex-direction: column !important;
        align-items: stretch !important;
    }
    .demo-shell > div,
    .evidence-pair > div {
        width: 100% !important;
        min-width: 0 !important;
        flex: 1 1 100% !important;
    }
}
@media (max-width: 720px) {
    .weld-crop-grid {
        grid-template-columns: 1fr;
    }
    .gradio-container .contain {
        padding-left: 10px;
        padding-right: 10px;
    }
    .app-topbar {
        position: relative;
        flex-direction: column !important;
        gap: 10px;
        margin: -16px -10px 0;
        padding: 14px 16px;
    }
    .app-topbar > div,
    .demo-shell > div,
    .evidence-pair > div {
        width: 100% !important;
        min-width: 0 !important;
        flex: 1 1 100% !important;
    }
    .brand-mark {
        width: 36px;
        min-width: 36px;
        min-height: 36px;
    }
    .brand-mark::before { inset: 9px; }
    .brand-mark::after { inset: 15px; }
    .brand-copy strong {
        font-size: 1.05rem;
    }
    .brand-copy small {
        display: none;
    }
    .status-chip textarea,
    .status-chip input {
        min-height: 36px !important;
    }
    .hero-panel {
        margin-top: 12px;
        padding: 17px;
    }
    .hero-panel h1 {
        font-size: 1.85rem;
    }
    .hero-copy {
        font-size: 0.86rem;
    }
    .runtime-control-panel,
    .custom-card,
    .dual-board-card {
        padding: 14px;
        border-radius: var(--radius-md);
    }
    .runtime-action-bar,
    .demo-shell,
    .evidence-pair {
        flex-direction: column !important;
        align-items: stretch !important;
    }
    .runtime-select-grid > div {
        min-width: 135px;
    }
    .board-head,
    .process-head {
        display: block;
    }
    .process-badge {
        margin-top: 10px;
    }
    .method-kpi-grid,
    .method-stage-grid {
        grid-template-columns: 1fr;
    }
    .prompt-card {
        min-height: 260px !important;
        max-height: 340px !important;
    }
    .logic-view-img,
    .evidence-image {
        height: 290px !important;
        min-height: 290px !important;
    }
}
@media (prefers-reduced-motion: reduce) {
    *, *::before, *::after {
        scroll-behavior: auto !important;
        transition-duration: 0.01ms !important;
        animation-duration: 0.01ms !important;
        animation-iteration-count: 1 !important;
    }
}

/* Compact research-workspace pass: the controls remain visible without competing
   with the evidence board, and Gradio's generated rows become real responsive grids. */
:root {
    --canvas: #F4F7F7;
    --panel: #FFFFFF;
    --panel-raised: #F8FBFB;
    --rule: #D9E3E4;
    --ink: #17242D;
    --muted: #66757C;
    --signal: #176F86;
    --signal-hover: #0D5C6F;
    --signal-subtle: #E7F2F4;
    --builder: #A96722;
    --builder-subtle: #F8F0E7;
    --pass: #19745A;
    --pass-subtle: #E9F5EF;
    --radius-sm: 6px;
    --radius-md: 8px;
    --radius-lg: 10px;
    --shadow-sm: 0 1px 2px rgba(23, 36, 45, 0.04);
    --shadow-md: 0 8px 20px rgba(23, 36, 45, 0.07);
}
html, body, .gradio-container {
    background:
        linear-gradient(90deg, rgba(207, 220, 220, 0.46) 1px, transparent 1px) 0 0 / 96px 96px,
        linear-gradient(0deg, rgba(207, 220, 220, 0.28) 1px, transparent 1px) 0 0 / 96px 96px,
        var(--canvas) !important;
}
.gradio-container .contain {
    width: 100% !important;
    max-width: none !important;
    padding-bottom: 28px;
}
.app-topbar {
    position: relative;
    z-index: 1;
    margin: -16px -16px 0;
    padding: 12px 24px;
    border-bottom: 1px solid #203F4B;
    background: #17303E;
    box-shadow: none;
    backdrop-filter: none;
}
.brand-block {
    gap: 10px;
}
.brand-mark {
    width: 34px;
    min-width: 34px;
    min-height: 34px;
    border-radius: 8px;
    border-color: rgba(150, 220, 217, 0.48);
    box-shadow: none;
}
.brand-mark::before { inset: 8px; }
.brand-mark::after { inset: 14px; }
.brand-copy strong {
    font-family: Inter, "IBM Plex Sans", "Segoe UI", "Microsoft YaHei UI", system-ui, sans-serif;
    font-size: 1rem;
    font-weight: 800;
}
.brand-copy small {
    font-size: 0.75rem;
    color: #B8CBD0;
}
.status-chip textarea,
.status-chip input {
    min-height: 34px !important;
    border-radius: var(--radius-sm) !important;
    border-color: rgba(184, 220, 220, 0.32) !important;
    background: rgba(255, 255, 255, 0.08) !important;
    box-shadow: none !important;
    font-size: 0.72rem !important;
    line-height: 1.25 !important;
    resize: none !important;
}
.hero-panel {
    margin: 16px 0 12px;
    padding: 18px 20px;
    border-radius: var(--radius-md);
    border-color: var(--rule);
    border-left: 4px solid var(--signal);
    background: var(--panel);
    box-shadow: none;
}
.hero-panel h1 {
    font-size: clamp(1.9rem, 3vw, 2.55rem);
    line-height: 1.06;
    letter-spacing: 0;
}
.eyebrow {
    color: var(--signal);
    font-size: 0.72rem;
    letter-spacing: 0;
}
.hero-copy {
    max-width: 760px;
    margin-bottom: 0;
}
.hero-metrics {
    gap: 6px;
}
.hero-metrics span {
    padding: 7px 10px;
    border-radius: var(--radius-sm);
    background: var(--panel-raised);
}
.runtime-control-panel {
    position: relative;
    top: auto;
    z-index: 1;
    margin: 12px 0 16px;
    padding: 16px 18px;
    border-radius: var(--radius-md);
    border-color: var(--rule);
    background: var(--panel);
    box-shadow: var(--shadow-sm);
    backdrop-filter: none;
}
.runtime-action-bar {
    display: grid !important;
    grid-template-columns: minmax(0, 1fr) minmax(190px, 240px);
    align-items: end;
    gap: 16px;
}
.runtime-action-bar > div {
    width: auto !important;
    min-width: 0 !important;
    flex: initial !important;
}
.runtime-heading .section-copy {
    max-width: 600px;
}
.runtime-action button,
.sample-actions button {
    min-height: 42px !important;
}
.runtime-select-grid {
    display: block !important;
    margin-top: 14px;
}
.runtime-select-grid > .form {
    display: grid !important;
    grid-template-columns: repeat(4, minmax(0, 1fr));
    width: 100% !important;
    min-width: 0 !important;
    gap: 12px;
    overflow: visible !important;
    border-color: transparent !important;
    background: transparent !important;
}
.runtime-select-grid > .form > .block {
    width: 100% !important;
    min-width: 0 !important;
    flex: initial !important;
}
.runtime-control-panel .block.padded,
.custom-card .block.padded {
    padding: 0 !important;
    border: 0 !important;
    overflow: visible !important;
}
.runtime-control-panel [data-testid="block-info"],
.custom-card [data-testid="block-info"] {
    display: block !important;
    margin: 0 0 6px !important;
    padding: 0 !important;
    border-radius: 0 !important;
    background: transparent !important;
    color: var(--muted) !important;
    font-size: 0.73rem !important;
    font-weight: 800 !important;
}
.runtime-control-panel .block > .container > .wrap,
.custom-card .block > .container > .wrap {
    border: 1px solid var(--rule) !important;
    border-radius: var(--radius-sm) !important;
    background: #FFFFFF !important;
    box-shadow: 0 1px 2px rgba(23, 36, 45, 0.03) !important;
}
.runtime-control-panel .block > .container > .wrap:focus-within,
.custom-card .block > .container > .wrap:focus-within {
    border-color: var(--signal) !important;
    box-shadow: 0 0 0 3px rgba(23, 111, 134, 0.14) !important;
}
.runtime-control-panel details {
    margin-top: 12px;
    border-top: 1px solid var(--rule);
}
.runtime-control-panel details summary {
    min-height: 42px;
    padding: 11px 2px;
    color: var(--ink);
    font-size: 0.88rem;
    font-weight: 800;
}
.runtime-setup {
    margin-bottom: 0 !important;
}
.section-title {
    border-left-width: 3px;
    padding-left: 9px;
    font-size: 1rem;
}
.section-copy {
    font-size: 0.86rem;
    line-height: 1.45;
}
.custom-card,
.dual-board-card {
    padding: 16px;
    border-radius: var(--radius-md);
    border-color: var(--rule);
    box-shadow: var(--shadow-sm);
}
.demo-shell {
    gap: 16px;
}
.sample-actions {
    display: grid !important;
    grid-template-columns: minmax(0, 1fr) minmax(0, 1fr);
    gap: 8px;
}
.sample-actions button {
    width: 100% !important;
    margin-top: 0;
}
.desk-btn-primary {
    background: var(--signal) !important;
    border-color: var(--signal-hover) !important;
    box-shadow: none !important;
    transition: background 0.16s ease !important;
}
.desk-btn-primary:hover {
    background: var(--signal-hover) !important;
    transform: none;
    box-shadow: none !important;
}
.desk-btn-secondary {
    background: #FFFFFF !important;
    border-color: var(--rule) !important;
}
.gradio-container button {
    border-radius: var(--radius-sm) !important;
}
.gradio-container button:focus-visible,
.gradio-container input:focus-visible,
.gradio-container textarea:focus-visible,
.gradio-container select:focus-visible,
.gradio-container summary:focus-visible {
    outline-color: rgba(23, 111, 134, 0.32) !important;
}
.weld-reference-card {
    margin: 0;
    padding: 10px;
    border: 1px solid #C5DDD5;
    border-radius: var(--radius-md);
    background: #FFFFFF;
}
.weld-reference-card img {
    display: block;
    width: 100%;
    max-height: 190px;
    object-fit: contain;
    border-radius: var(--radius-sm);
    background: var(--panel-raised);
}
.weld-reference-card figcaption {
    margin-top: 7px;
    color: var(--ink);
    font-size: 0.75rem;
    font-weight: 800;
    text-align: center;
}
@media (max-width: 900px) {
    .runtime-action-bar {
        grid-template-columns: minmax(0, 1fr) 210px;
    }
    .runtime-select-grid > .form {
        grid-template-columns: repeat(2, minmax(0, 1fr));
    }
}
@media (max-width: 680px) {
    .app-topbar {
        margin: -16px -10px 0;
        padding: 11px 14px;
    }
    .hero-panel {
        margin-top: 12px;
        padding: 16px;
    }
    .hero-panel h1 {
        font-size: 1.8rem;
    }
    .runtime-control-panel,
    .custom-card,
    .dual-board-card {
        padding: 14px;
    }
    .runtime-action-bar {
        grid-template-columns: 1fr;
        gap: 10px;
    }
    .runtime-action button {
        width: 100% !important;
    }
    .hero-metrics {
        grid-template-columns: repeat(3, minmax(0, 1fr));
        gap: 5px;
    }
    .hero-metrics span {
        display: grid;
        gap: 2px;
        padding: 7px 8px;
        font-size: 0.76rem;
    }
}
@media (max-width: 500px) {
    .runtime-select-grid > .form,
    .sample-actions {
        grid-template-columns: 1fr;
    }
    .hero-metrics {
        grid-template-columns: 1fr;
    }
}
"""

def create_ui(initial_dataset=None):
    initial_dataset = str(initial_dataset or os.environ.get("GLLS_INITIAL_DATASET", "mvtec")).lower()
    if initial_dataset not in {"mvtec", "visa", WELD_DATASET}:
        initial_dataset = "mvtec"
    desk_theme = gr.themes.Soft(primary_hue="blue", neutral_hue="slate").set(
        block_radius="8px",
        button_primary_background_fill="#176F86",
        button_primary_background_fill_hover="#0D5C6F",
    )

    MVTEC_CLASSES = [
        "bottle", "cable", "capsule", "carpet", "grid", "hazelnut", "leather", 
        "metal_nut", "pill", "screw", "tile", "toothbrush", "transistor", "wood", "zipper"
    ]
    VISA_CLASSES = [
        "candle", "capsules", "cashew", "chewinggum", "fryum", "macaroni1", 
        "macaroni2", "pcb1", "pcb2", "pcb3", "pcb4", "pipe_fryum"
    ]
    dataset_classes = {
        "mvtec": MVTEC_CLASSES,
        "visa": VISA_CLASSES,
        WELD_DATASET: [WELD_CATEGORY],
    }
    initial_classes = dataset_classes[initial_dataset]
    initial_category = initial_classes[0]
    initial_localizer_choices = (
        ["AdaptCLIP"]
        if initial_dataset == WELD_DATASET
        else ["Auto", "ABounD", "AdaptCLIP"]
    )
    initial_localizer = "AdaptCLIP" if initial_dataset == WELD_DATASET else "Auto"
    initial_shot_choices = ["1"] if initial_dataset == WELD_DATASET else ["1", "0"]
    initial_checkpoint = _default_adaptclip_checkpoint(initial_dataset)
    
    with gr.Blocks(css=CSS, theme=desk_theme, title="GLLS Method Desk") as demo:

        with gr.Row(elem_classes="app-topbar"):
            with gr.Column(scale=3):
                gr.HTML(
                    """
                    <div class="brand-block">
                      <div class="brand-mark"></div>
                      <div class="brand-copy">
                        <strong>GLLS Method Desk</strong>
                        <small>Global logic, local evidence, SAM3 refinement, and PVLA/RAG provenance.</small>
                      </div>
                    </div>
                    """
                )
            with gr.Column(scale=2, elem_classes="status-chip"):
                status_box = gr.Textbox(label="Runtime", value="Not initialized", interactive=False, container=False)

        gr.HTML(
            """
            <section class="hero-panel">
              <div>
                <div class="eyebrow">Dual-stream anomaly reasoning demo</div>
                <h1>GLLS Verification Board</h1>
                <p class="hero-copy">
                  Load a review sample, initialize the selected weights, and inspect the evidence
                  passed from global logic and local search to the final decision.
                </p>
              </div>
              <div class="hero-metrics">
                <span><b>1</b><em>Configure runtime</em></span>
                <span><b>2</b><em>Load review sample</em></span>
                <span><b>3</b><em>Run verifier</em></span>
              </div>
            </section>
            """
        )

        with gr.Column(elem_classes="runtime-control-panel"):
            with gr.Row(elem_classes="runtime-action-bar"):
                with gr.Column(scale=3, elem_classes="runtime-heading"):
                    gr.Markdown("### Runtime", elem_classes="section-title")
                    gr.Markdown(
                        "Initialize the selected weights before running a review sample.",
                        elem_classes="section-copy",
                    )
                with gr.Column(scale=1, elem_classes="runtime-action"):
                    btn_init = gr.Button("Initialize runtime", elem_classes="desk-btn-primary")
            with gr.Row(elem_classes="runtime-select-grid"):
                dd_gpu = gr.Dropdown(_gpu_choices(), value=_default_gpu_id(), label="GPU")
                dd_dataset = gr.Dropdown(
                    ["mvtec", "visa", WELD_DATASET],
                    label="Dataset",
                    value=initial_dataset,
                )
                dd_localizer = gr.Dropdown(
                    initial_localizer_choices,
                    label="Localizer",
                    value=initial_localizer,
                )
                dd_kshot = gr.Dropdown(
                    initial_shot_choices,
                    value="1",
                    label="Shot",
                )
            with gr.Accordion("Selected adapted weights", open=False):
                txt_runtime_weights = gr.Textbox(
                    label="Resolved paths",
                    value=_runtime_weight_summary(
                        _runtime_weight_config(initial_dataset, initial_localizer, "1", initial_checkpoint)
                    ),
                    lines=5,
                    interactive=False,
                )
            with gr.Accordion("Advanced runtime paths", open=False, elem_classes="runtime-setup"):
                with gr.Row():
                    dd_model_type = gr.Dropdown(
                        ["qwen3-vl", "qwen2.5-vl", "llava-onevision"],
                        label="VLM architecture",
                        value="qwen3-vl",
                    )
                    cb_vllm = gr.Checkbox(label="Use vLLM if installed", value=False, interactive=True)
                    sl_gpu_util = gr.Slider(0.3, 0.95, value=0.85, step=0.05, label="vLLM memory fraction")
                with gr.Row():
                    p_model = gr.Textbox(DEFAULT_VLM_MODEL_PATH, label="VLM model folder")
                    p_sam = gr.Textbox(DEFAULT_SAM3_PATH, label="SAM3 checkpoint")
                with gr.Row():
                    p_data = gr.Textbox(DEFAULT_MMAD_ROOT, label="MMAD dataset root")
                    p_qa = gr.Textbox(DEFAULT_QA_ROOT, label="Curated QA root")
                with gr.Row():
                    p_ckpt = gr.Textbox(
                        initial_checkpoint,
                        label="AdaptCLIP checkpoint (used when AdaptCLIP is selected)",
                    )
                    p_graph = gr.Textbox(DEFAULT_GRAPH_ROOT, label="PVLA graph cache root")

        with gr.Row(elem_classes="demo-shell"):
            with gr.Column(scale=1, min_width=390):
                with gr.Column(elem_classes="custom-card"):
                    gr.Markdown("### Review Sample", elem_classes="section-title")
                    gr.Markdown("Choose a category and review case, then run the verifier.", elem_classes="section-copy")
                    with gr.Row():
                        dd_cat = gr.Dropdown(initial_classes, label="Category", value=initial_category, interactive=True)
                    with gr.Row(elem_classes="sample-actions"):
                        btn_load = gr.Button("Load cases", elem_classes="desk-btn-secondary")
                        btn_run_main = gr.Button("Run GLLS", variant="primary", elem_classes="desk-btn-primary")
                    with gr.Row():
                        dd_sub_type = gr.Dropdown(label="Defect folder", choices=["All"], value="All")
                        dd_logic = gr.Dropdown(label="Task", choices=["All"], value="All")
                    txt_search_case = gr.Textbox(placeholder="image name / question / task", label="Search")
                    dd_samples_list = gr.Dropdown(label="Review case / Question", choices=[], interactive=True)
                    img_preview_in = gr.Image(label="Input image", type="pil", height=300, elem_classes="desk-image-upload")
                    txt_q_in = gr.Textbox(label="Question", lines=2)
                    with gr.Accordion("Answer options", open=False):
                        txt_gt_out = gr.Textbox(label="Ground truth", interactive=False)
                        txt_opts_json = gr.Code(label="Options", language="json")

            with gr.Column(scale=2):
                with gr.Column(elem_classes="dual-board-card"):
                    gr.Markdown("### Dual-Stream Verification Board", elem_classes="section-title")
                    md_method_process = gr.HTML(
                        render_weld_placeholder()
                        if initial_dataset == WELD_DATASET
                        else _method_process_placeholder()
                    )

                    with gr.Accordion("PVLA details / RAG recall", open=False):
                        with gr.Row(elem_classes="evidence-pair"):
                            with gr.Column(scale=1):
                                gal_rag_source = gr.Gallery(
                                    label="RAG visual references",
                                    show_label=True,
                                    columns=4,
                                    rows=1,
                                    height=190,
                                    object_fit="contain",
                                    type="filepath",
                                )
                            with gr.Column(scale=1, elem_classes="evidence-detail"):
                                txt_rag_manual = gr.TextArea(label="Retrieved text knowledge", lines=7, interactive=False)

                    with gr.Accordion("Run visuals — logic view, heatmap & crops", open=False):
                        gr.Markdown("Stream 1 · Global & logic", elem_classes="section-title-sm")
                        with gr.Row(equal_height=True, elem_classes="evidence-pair"):
                            with gr.Column(scale=1):
                                img_logic_debug = gr.Image(label="Logic view", type="pil", elem_classes="logic-view-img", show_label=True, interactive=False)
                            with gr.Column(scale=1, elem_classes="evidence-detail"):
                                md_p1_prompt = gr.Markdown(value="Waiting for a run.", elem_classes="prompt-card")
                        gr.Markdown("Stream 2 · Fine-grained & actions", elem_classes="section-title-sm")
                        with gr.Row(elem_classes="evidence-pair"):
                            with gr.Column(scale=1):
                                img_hm_out = gr.Image(
                                    label="Anomaly heatmap",
                                    type="pil",
                                    height=250,
                                    elem_classes="evidence-image",
                                    interactive=False,
                                )
                            with gr.Column(scale=1):
                                state_redbox_pil = gr.Image(label="Global red-box trace", type="pil", height=250, elem_classes="evidence-image", interactive=False)
                        gal_sam3_preview = gr.Gallery(
                            label="Top-K crops and local views",
                            columns=4,
                            height=230,
                            object_fit="contain",
                            preview=True,
                            interactive=False,
                        )
                        file_sam3_editor = gr.File(
                            label="Selected focus files",
                            file_count="multiple",
                            type="filepath",
                            visible=False,
                        )

                    with gr.Accordion("Final Verdict", open=True):
                        md_final_res = gr.Markdown("### Waiting for a run")

                    with gr.Accordion("Developer trace / Export", open=False):
                        with gr.Tabs():
                            with gr.TabItem("Prompts"):
                                txt_cot_edit = gr.TextArea(label="Final prompt", lines=8, interactive=False)
                            with gr.TabItem("Run logs"):
                                txt_logs_stream = gr.TextArea(elem_classes="log-box", lines=12, show_copy_button=True)
                            with gr.TabItem("Export"):
                                cb_export_docx = gr.Checkbox(
                                    label="Include Word report (.docx) inside ZIP",
                                    value=True,
                                    interactive=True,
                                )
                                if hasattr(gr, "DownloadButton"):
                                    try:
                                        btn_export_zip = gr.DownloadButton(
                                            "Download full run bundle",
                                            variant="primary",
                                            size="lg",
                                            elem_classes="desk-btn-primary",
                                        )
                                    except TypeError:
                                        btn_export_zip = gr.DownloadButton("Download full run bundle")
                                    file_export_zip = None
                                else:
                                    btn_export_zip = gr.Button(
                                        "Generate export bundle",
                                        variant="primary",
                                        size="lg",
                                        elem_classes="desk-btn-primary",
                                    )
                                    file_export_zip = gr.File(label="Export ZIP", interactive=False, type="filepath")
                                txt_export_status = gr.Textbox(label="Export status", interactive=False)

        # --- States ---
        txt_hidden_t_type = gr.Textbox(visible=False)
        state_crops_paths = gr.State()

        # --- Event Bindings ---
        
        # 1. Initialize
        btn_init.click(
            lambda pm, pl, pg, ps, pd, pq, gpu, mt, vlm, gu, dn, ks, lc: global_sys.initialize(
                pm, pl, pg, ps, pd, pq, gpu, mt, vlm, gu, dn, ks, lc
            ),
            inputs=[
                p_model,
                p_ckpt,
                p_graph,
                p_sam,
                p_data,
                p_qa,
                dd_gpu,
                dd_model_type,
                cb_vllm,
                sl_gpu_util,
                dd_dataset,
                dd_kshot,
                dd_localizer,
            ],
            outputs=status_box,
            api_name=False,
            concurrency_limit=1,
            concurrency_id="glls_runtime",
        )
        
        # 2. Update Category List when Dataset Changes
        def update_cat_list(ds_name):
            if ds_name == "mvtec":
                ckpt = _default_adaptclip_checkpoint("mvtec")
                categories = MVTEC_CLASSES
                localizer_update = gr.update(choices=["Auto", "ABounD", "AdaptCLIP"], value="Auto")
                shot_update = gr.update(choices=["1", "0"], value="1")
                localizer_value = "Auto"
            elif ds_name == "visa":
                ckpt = _default_adaptclip_checkpoint("visa")
                categories = VISA_CLASSES
                localizer_update = gr.update(choices=["Auto", "ABounD", "AdaptCLIP"], value="Auto")
                shot_update = gr.update(choices=["1", "0"], value="1")
                localizer_value = "Auto"
            else:
                ckpt = _default_adaptclip_checkpoint(WELD_DATASET)
                categories = [WELD_CATEGORY]
                localizer_update = gr.update(choices=["AdaptCLIP"], value="AdaptCLIP")
                shot_update = gr.update(choices=["1"], value="1")
                localizer_value = "AdaptCLIP"
            return (
                gr.update(choices=categories, value=categories[0]),
                gr.update(value=ckpt),
                localizer_update,
                shot_update,
                ckpt,
                categories[0],
                localizer_value,
            )

        def runtime_preview(ds_name, localizer_choice, k_shot, ckpt_path, gpu_id):
            try:
                cfg = _runtime_weight_config(ds_name, localizer_choice, k_shot, ckpt_path)
            except Exception as e:
                return f"Invalid runtime configuration: {e}", f"Runtime config invalid: {e}"
            ok, _, _ = global_sys.runtime_match_status(ds_name, localizer_choice, k_shot, ckpt_path, gpu_id)
            if ok:
                status_msg = "Runtime ready for selected weights."
            elif global_sys.is_initialized:
                status_msg = "Runtime stale: initialize to load the selected weights."
            else:
                status_msg = "Runtime not initialized for selected weights."
            return _runtime_weight_summary(cfg), status_msg
        
        def empty_case_payload():
            return None, "", "", "", "{}"

        def sample_payload_from_choice(idx_str):
            if str(idx_str or "").startswith(WELD_TASK_PREFIX):
                payload = global_sys.weld_runtime.sample_payload(idx_str)
                return (
                    gr.update(value=payload["image"], interactive=False),
                    payload["question"],
                    payload["task"],
                    payload["ground_truth"],
                    json.dumps(payload["options"], indent=2, ensure_ascii=False),
                )
            if not idx_str or not global_sys.data_manager:
                return empty_case_payload()
            idx = int(idx_str.split("|")[0].strip())
            img, q, t, gt, opts = global_sys.data_manager.get_sample_by_idx(idx)
            return gr.update(value=img, interactive=True), q, t, gt, json.dumps(opts, indent=2, ensure_ascii=False)

        # 3. Load Dataset & Reset Subclass/Logic Dropdowns
        def on_cat_load(ds_name, c, localizer_choice="Auto", k_shot=1, ckpt_path=DEFAULT_ADAPTCLIP_CHECKPOINT_PATH, gpu_id="0"):
            if ds_name == WELD_DATASET:
                weight_summary, runtime_msg = runtime_preview(
                    ds_name,
                    "AdaptCLIP",
                    "1",
                    ckpt_path,
                    gpu_id,
                )
                try:
                    choices = global_sys.weld_runtime.sample_choices()
                    selected_sample = choices[0][1] if choices else None
                    msg = (
                        f"[WELD] {WELD_CATEGORY} loaded: {len(choices)} qualitative review images | "
                        f"{runtime_msg}"
                    )
                except Exception as e:
                    choices = []
                    selected_sample = None
                    msg = f"Weld demo unavailable: {e} | {runtime_msg}"
                return (
                    gr.update(
                        choices=["All", "annotated_anomaly", "allowed_context", "unlabeled"],
                        value="All",
                    ),
                    gr.update(choices=["All", "weld_inspection"], value="All"),
                    gr.update(choices=choices, value=selected_sample),
                    *sample_payload_from_choice(selected_sample),
                    weight_summary,
                    msg,
                    render_weld_placeholder(),
                )
            if not global_sys.data_manager:
                global_sys.data_manager = DatasetManager(DEFAULT_MMAD_ROOT, DEFAULT_QA_ROOT, ds_name or "mvtec")
            if ds_name and global_sys.data_manager.dataset_name != ds_name:
                global_sys.data_manager = DatasetManager(
                    global_sys.data_manager.root_path,
                    global_sys.data_manager.qa_root_path,
                    ds_name,
                )
            msg = global_sys.data_manager.load_subclass(c)
            weight_summary, runtime_msg = runtime_preview(ds_name, localizer_choice, k_shot, ckpt_path, gpu_id)
            msg = f"{msg} | {runtime_msg}"
            # FORCE RESET of dropdowns to prevent cross-dataset sticky values
            subs = ["All"] + global_sys.data_manager.subfolder_choices
            logics = ["All"] + global_sys.data_manager.task_type_choices
            sample_choices = global_sys.data_manager.filter_samples()
            selected_sample = sample_choices[0] if sample_choices else None
            return (
                gr.update(choices=subs, value="All"),
                gr.update(choices=logics, value="All"),
                gr.update(choices=sample_choices, value=selected_sample),
                *sample_payload_from_choice(selected_sample),
                weight_summary,
                msg,
                render_dual_stream_placeholder(c),
            )

        qa_load_outputs = [
            dd_sub_type,
            dd_logic,
            dd_samples_list,
            img_preview_in,
            txt_q_in,
            txt_hidden_t_type,
            txt_gt_out,
            txt_opts_json,
            txt_runtime_weights,
            status_box,
            md_method_process,
        ]

        def on_dataset_change(ds_name, gpu_id):
            (
                cat_update,
                ckpt_update,
                localizer_update,
                shot_update,
                next_ckpt,
                next_category,
                next_localizer,
            ) = update_cat_list(ds_name)
            return (
                cat_update,
                ckpt_update,
                localizer_update,
                shot_update,
                *on_cat_load(ds_name, next_category, next_localizer, "1", next_ckpt, gpu_id),
            )

        dd_dataset.change(
            on_dataset_change,
            [dd_dataset, dd_gpu],
            [dd_cat, p_ckpt, dd_localizer, dd_kshot, *qa_load_outputs],
            api_name=False,
            queue=False,
        )
        dd_cat.change(
            on_cat_load,
            [dd_dataset, dd_cat, dd_localizer, dd_kshot, p_ckpt, dd_gpu],
            qa_load_outputs,
            api_name=False,
            queue=False,
        )
        btn_load.click(
            on_cat_load,
            [dd_dataset, dd_cat, dd_localizer, dd_kshot, p_ckpt, dd_gpu],
            qa_load_outputs,
            api_name=False,
            queue=False,
        )
        demo.load(
            on_cat_load,
            [dd_dataset, dd_cat, dd_localizer, dd_kshot, p_ckpt, dd_gpu],
            qa_load_outputs,
            api_name=False,
        )

        for runtime_trigger in [dd_localizer, dd_kshot, p_ckpt, dd_gpu]:
            runtime_trigger.change(
                runtime_preview,
                [dd_dataset, dd_localizer, dd_kshot, p_ckpt, dd_gpu],
                [txt_runtime_weights, status_box],
                api_name=False,
                queue=False,
            )

        # 4. Filter Samples
        def on_filter_update(ds_name, s_f, l_f, search_t):
            if ds_name == WELD_DATASET:
                try:
                    choices = global_sys.weld_runtime.sample_choices(s_f, search_t)
                    selected = choices[0][1] if choices else None
                except Exception:
                    choices = []
                    selected = None
                return (
                    gr.update(choices=choices, value=selected),
                    *sample_payload_from_choice(selected),
                )
            if not global_sys.data_manager:
                return gr.update(choices=[]), *empty_case_payload()
            choices = global_sys.data_manager.filter_samples(
                subfolder=None if s_f == "All" else s_f,
                task_type=None if l_f == "All" else l_f,
                search_query=search_t
            )
            selected = choices[0] if choices else None
            return (
                gr.update(choices=choices, value=selected),
                *sample_payload_from_choice(selected),
            )

        for trigger in [dd_sub_type, dd_logic, txt_search_case]:
            trigger.change(
                on_filter_update,
                [dd_dataset, dd_sub_type, dd_logic, txt_search_case],
                [dd_samples_list, img_preview_in, txt_q_in, txt_hidden_t_type, txt_gt_out, txt_opts_json],
                api_name=False,
                queue=False,
            )

        # 5. Select Case & Update Options State
        def on_select_case_id(idx_str):
            return sample_payload_from_choice(idx_str)

        dd_samples_list.change(
            on_select_case_id, dd_samples_list, 
            [img_preview_in, txt_q_in, txt_hidden_t_type, txt_gt_out, txt_opts_json],
            api_name=False,
            queue=False,
        )

        # 6. Run Actions
        btn_run_main.click(
            search_runner_wrapper,
            inputs=[
                img_preview_in,
                txt_q_in,
                dd_cat,
                txt_hidden_t_type,
                txt_opts_json,
                txt_gt_out,
                dd_dataset,
                dd_localizer,
                dd_kshot,
                p_ckpt,
                dd_gpu,
            ],
            outputs=[gal_rag_source, txt_rag_manual, img_hm_out, gal_sam3_preview, file_sam3_editor, txt_logs_stream, md_final_res, txt_cot_edit, state_redbox_pil, state_crops_paths, img_logic_debug, md_p1_prompt, md_method_process],
            api_name=False,
            concurrency_limit=1,
            concurrency_id="glls_runtime",
        )

        export_outputs = [file_export_zip, txt_export_status] if file_export_zip is not None else [btn_export_zip, txt_export_status]
        btn_export_zip.click(
            export_current_run_bundle,
            inputs=[
                img_preview_in,
                txt_q_in,
                txt_gt_out,
                dd_cat,
                txt_hidden_t_type,
                txt_opts_json,
                gal_rag_source,
                txt_rag_manual,
                img_hm_out,
                gal_sam3_preview,
                file_sam3_editor,
                md_final_res,
                txt_cot_edit,
                img_logic_debug,
                md_p1_prompt,
                txt_logs_stream,
                cb_export_docx,
            ],
            outputs=export_outputs,
            api_name=False,
        )

    return demo

def main():
    import argparse

    parser = argparse.ArgumentParser(description="Launch the GLLS Gradio visualization UI.")
    parser.add_argument("--host", default="0.0.0.0", help="Host/interface to bind.")
    parser.add_argument("--port", type=int, default=None, help="Port to bind. Defaults to the first free port from --start_port.")
    parser.add_argument("--start_port", type=int, default=7860, help="First port to try when --port is not set.")
    parser.add_argument("--max_queue", type=int, default=5, help="Gradio queue max size.")
    parser.add_argument("--show_api", action="store_true", help="Show Gradio API docs.")
    parser.add_argument(
        "--initial_dataset",
        choices=["mvtec", "visa", WELD_DATASET],
        default=os.environ.get("GLLS_INITIAL_DATASET", "mvtec"),
        help="Dataset selected when the shared demo first opens.",
    )
    args = parser.parse_args()

    # Gradio uses httpx for local connectivity checks. If http_proxy or
    # https_proxy is set, localhost requests may be routed through the proxy and
    # fail with httpx.ConnectError. Keep loopback traffic out of the proxy.
    _no_proxy_hosts = "localhost,127.0.0.1,0.0.0.0"
    os.environ["NO_PROXY"] = ",".join(
        [h for h in (os.environ.get("NO_PROXY", ""), _no_proxy_hosts) if h]
    )
    os.environ["no_proxy"] = ",".join(
        [h for h in (os.environ.get("no_proxy", ""), _no_proxy_hosts) if h]
    )

    port = args.port if args.port is not None else find_free_port(args.start_port)
    print(f"GLLS Method Desk launched at http://localhost:{port}")
    patch_gradio_compat()
    app = create_ui(args.initial_dataset)
    app.queue(max_size=args.max_queue).launch(server_name=args.host, server_port=port, show_api=args.show_api)


if __name__ == "__main__":
    main()
